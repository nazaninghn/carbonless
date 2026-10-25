"""Stage 3 audit: scenarios on the representative inventory (see test_stage3_oracle.py).

* electricity, location-based vs market-based: supplier declaration (right market / wrong market / wrong year)
  and facility-specific approved factors (Ozel Talep), each recalculated here by hand;
* unapproved data (a pending change, a pending new entry, a rejected entry) reaches no surface: not the API,
  CSV, Excel, nor either PDF.

The inventory: Gebze 50,000 kWh (distribution 0.469) + Depo 20,000 kWh (transmission 0.436), location-based
Scope 2 = 23,450 + 8,720 = 32,170 kg; inventory total 396,315.3 kg."""
import csv
import io
from decimal import Decimal as D

from django.contrib.auth.models import User
from django.core.cache import cache
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from companies.models import CompanyMembership
from emissions.models import CustomEmissionRequest, EmissionEntry, EmissionFactor
from . import tests as _base
from . import test_stage3_oracle as _o
from .models import ReportStep
from .step_entries import resync_inventory, scope2_market_based

LOCATION_S2 = 32170.0
TOTAL = 396315.3


class _Inventory(TestCase):
    setUp = _o.RepresentativeInventoryTests.setUp
    _summary, _entries, _pdf_text = (_o.RepresentativeInventoryTests._summary, _o.RepresentativeInventoryTests._entries,
                                     _o.RepresentativeInventoryTests._pdf_text)

    def _step(self, step_id, answer):
        cache.clear()
        return _base.StepEntriesTests._step(self, step_id, answer)

    def _market(self):
        answers = dict(ReportStep.objects.filter(report=self.report).values_list('step_id', 'answer'))
        return scope2_market_based(answers, 2025, self.company)

    def _supplier(self, value='0.3', market='TR', year='2025', unit='kgCO2e_kWh'):
        self._step('4A-EF', 'yes')
        self._step('4A-EF-a', {'ef_value': value, 'ef_unit': unit, 'ef_source': 'Supplier X', 'ef_year': year,
                               'ef_market': market})

    def _request(self, facility, method, factor):
        return CustomEmissionRequest.objects.create(
            user=self.user, company=self.company, scope='scope2', category_name='Electricity', source_name='Grid',
            description='', unit='kWh', quantity=D('1'), year=2025, month=1, facility=facility, status='approved',
            electricity_method=method, approved_factor_kg_co2e=D(factor), calculated_co2e_kg=D(factor))


class ElectricityMethodTests(_Inventory):

    def test_without_a_declaration_market_based_equals_location_based(self):
        m = self._market()
        self.assertAlmostEqual(m['total_kg'], LOCATION_S2, places=2)
        self.assertEqual({r['basis'] for r in m['rows']}, {'location_fallback'})

    def test_a_supplier_declaration_for_the_right_market_and_year_sets_the_market_based_figure_only(self):
        self._supplier('0.3', 'TR', '2025')
        m = self._market()
        self.assertAlmostEqual(m['total_kg'], 70000 * 0.3, places=2)                      # 21,000 kg, by hand
        self.assertEqual({r['basis'] for r in m['rows']}, {'supplier'})
        s = self._summary()
        self.assertAlmostEqual(s['scope2_kg'], LOCATION_S2, places=2)                     # location-based unchanged
        self.assertAlmostEqual(s['total_kg'], TOTAL, places=2)                            # and never added to any total
        text = self._pdf_text(f'/api/questionnaire/{self.report.id}/iso-report/?lang=en')
        self.assertIn('21.00 t', text)                                                    # the market-based result is reported

    def test_a_declaration_for_another_market_is_not_applied_to_these_sites(self):
        self._supplier('0.3', 'DE', '2025')
        m = self._market()
        self.assertAlmostEqual(m['total_kg'], LOCATION_S2, places=2)                      # fell back to the grid, not 0.3
        self.assertTrue(m['outside_market'])

    def test_a_declaration_of_another_year_is_not_applied(self):
        self._supplier('0.3', 'TR', '2024')
        self.assertAlmostEqual(self._market()['total_kg'], LOCATION_S2, places=2)

    def test_a_declaration_in_mwh_is_scaled_to_kwh(self):
        self._supplier('300', 'TR', '2025', unit='kgCO2e_MWh')                             # 300 kg/MWh = 0.3 kg/kWh
        self.assertAlmostEqual(self._market()['total_kg'], 21000.0, places=2)

    def test_a_supplier_factor_never_replaces_the_location_based_grid_factor(self):
        self._supplier('0.1', 'TR', '2025')
        elec = [e for e in self._entries() if e.emission_factor.category == 'electricity']
        self.assertEqual(sorted(float(e.calculated_co2e_kg) for e in elec), [8720.0, 23450.0])

    def test_a_facility_market_based_request_wins_over_the_declaration_for_that_facility(self):
        self._supplier('0.3', 'TR', '2025')
        gebze = self.company.facilities.get(name='Gebze')
        self._request(gebze, 'market_based', '0.2')
        m = self._market()
        by_site = {r['site']: (r['basis'], r['kg']) for r in m['rows']}
        self.assertEqual(by_site['Gebze'], ('facility_request', 50000 * 0.2))             # 10,000
        self.assertEqual(by_site['Depo'][0], 'supplier')
        self.assertAlmostEqual(m['total_kg'], 10000 + 20000 * 0.3, places=2)              # 16,000
        self.assertAlmostEqual(self._summary()['total_kg'], TOTAL, places=2)

    def test_an_approved_location_based_request_replaces_that_facilitys_grid_factor_in_the_inventory(self):
        gebze = self.company.facilities.get(name='Gebze')
        self._request(gebze, 'location_based', '0.5')
        resync_inventory(self.company, 2025, self.user)                                    # what the admin save does
        elec = sorted(float(e.calculated_co2e_kg) for e in self._entries() if e.emission_factor.category == 'electricity'
                      or e.emission_factor.slug == 'calculated-scope2-electricity')
        self.assertEqual(elec, [8720.0, 25000.0])                                          # Gebze 50,000 x 0.5; Depo untouched
        self.assertAlmostEqual(self._summary()['total_kg'], TOTAL - 23450 + 25000, places=2)
        # the request's own amount is not added on top (it would count the same kWh twice)
        self.assertAlmostEqual(self._summary()['custom_emissions']['approved_total_kg'], 0.0, places=2)


class UnapprovedNeverCountsTests(_Inventory):

    def setUp(self):
        super().setUp()
        self.clerk = User.objects.create_user('clerk3', 'clerk3@test.com', 'pass12345')
        CompanyMembership.objects.create(user=self.clerk, company=self.company, role='data_entry')
        c = APIClient()
        c.force_authenticate(user=self.clerk)
        cache.clear()
        # a data-entry member's change to the approved electricity answer: waits for approval
        r = c.patch(f'/api/questionnaire/{self.report.id}/step/',
                    {'step': '4A-1', 'data': {'answer': {'1': '99999 kWh', '2': '20000 kWh'}}, 'language': 'en'}, format='json')
        self.assertEqual(r.status_code, 200, r.content)
        self.assertTrue(r.data['calc_feedback']['pending'])
        gas = EmissionFactor.objects.get(slug='natural-gas-m3', country='global')
        mk = lambda status, **kw: EmissionEntry.objects.create(
            user=self.owner if hasattr(self, 'owner') else self.user, company=self.company, emission_factor=gas, year=2025,
            month=3, quantity=D('1000'), status=status, description='form entry', **kw)
        mk('draft', rejected_reason='no invoice')            # rejected
        mk('submitted')                                       # a new entry waiting

    def test_the_unapproved_rows_really_exist(self):
        """Premise of the tests below: they are in the database, with real amounts, just not approved."""
        rows = EmissionEntry.objects.filter(company=self.company, year=2025).exclude(status='approved')
        self.assertEqual(sorted(r.status for r in rows), ['draft', 'submitted', 'submitted', 'submitted'])   # rejected; 2 rows of the pending change (one per facility); 1 new
        self.assertGreater(sum(float(r.calculated_co2e_kg) for r in rows), 40000)                  # incl. 99,999 kWh x 0.469

    def test_no_surface_counts_pending_or_rejected_data(self):
        s = self._summary()
        self.assertAlmostEqual(s['total_kg'], TOTAL, places=2)
        self.assertEqual(s['not_counted']['pending']['count'] + s['not_counted']['changes']['count'] > 0, True)
        self.assertEqual(s['not_counted']['rejected']['count'], 1)
        for lang in ('en', 'tr'):
            rows = list(csv.reader(io.StringIO(self.client.get(f'/api/emissions/export-csv/?year=2025&lang={lang}')
                                               .content.decode('utf-8').lstrip('﻿')), delimiter=';' if lang == 'tr' else ','))
            h = rows[0]
            ok = 'Onaylı' if lang == 'tr' else 'Approved'
            approved = [r for r in rows[1:] if r[h.index('Durum' if lang == 'tr' else 'Status')] == ok]
            cell = lambda v: float(str(v).replace(',', '.'))
            self.assertAlmostEqual(sum(cell(r[h.index('kg CO2e')]) for r in approved), TOTAL, places=1, msg=lang)
            self.assertGreater(len(rows) - 1, len(approved))          # the unapproved rows are listed, but marked
        xr = [list(r) for r in load_workbook(io.BytesIO(self.client.get('/api/emissions/export-excel/?year=2025&lang=en')
                                                        .content)).active.iter_rows(values_only=True)]
        xh = xr[0]
        self.assertAlmostEqual(sum(float(r[xh.index('kg CO2e')]) for r in xr[1:]
                                   if r and r[0] and r[xh.index('Status')] == 'Approved'), TOTAL, places=1)

    def test_neither_pdf_counts_them(self):
        pending_kg = 99999 * 0.469                                       # 46,899.53 kg if it had been counted
        emissions = self._pdf_text('/api/emissions/report/?year=2025&lang=en')
        self.assertIn('Total emissions: 396.32 tCO₂e', emissions)
        self.assertNotIn(f'{pending_kg:,.2f}', emissions)
        iso = self._pdf_text(f'/api/questionnaire/{self.report.id}/iso-report/?lang=en')
        self.assertIn('396.32 tonnes', iso)
        self.assertNotIn(f'{pending_kg / 1000:.2f}', iso)

    def test_questionnaire_total_is_the_approved_value_while_the_change_waits(self):
        from .step_entries import questionnaire_totals
        effective, approved = questionnaire_totals(self.report)
        self.assertAlmostEqual(approved, TOTAL, places=2)                # what the Wizard calls the approved questionnaire total
        self.assertGreater(effective, approved)                          # what it would be once approved: not what is reported


class ExtendedScope3Tests(TestCase):
    """Downstream transport, a leased asset and a PCAF investment, each calculated by hand."""

    def setUp(self):
        from django.core.management import call_command
        _base.StepEntriesTests.setUp(self)
        call_command('seed_factors', verbosity=0)
        call_command('seed_gas_splits', verbosity=0)
        for sid, ans in (
            ('K3C9-1', {'items': [{'transport_mode': 'TM-02', 'load_tonne': '2', 'distance_km': '50'}]}),
            ('K3C15-1', {'items': [{'asset_class': 'VA-01', 'investment_amount': '5000000', 'company_value': '40000000',
                                    'company_debt': '10000000', 'ghg_report_available': True,
                                    'company_emissions_tco2e': '1200'}]}),
            ('K3C8-1', {'items': [{'asset_type': 'KV-01', 'area_m2': '250', 'owner_declaration': True,
                                   'declaration_kwh': '12000', 'declaration_energy': 'electricity',
                                   'in_scope12': 'not_included'}]}),
        ):
            cache.clear()
            _base.StepEntriesTests._step(self, sid, ans)

    def _by_step(self):
        out = {}
        for e in EmissionEntry.objects.filter(company=self.company, year=2025).select_related('emission_factor'):
            out.setdefault(e.description.split(' ')[2], []).append(
                (e.emission_factor.slug, float(e.quantity), round(float(e.calculated_co2e_kg), 4)))
        return out

    def test_downstream_transport_is_tonne_km_times_the_factor(self):
        # 2 t x 50 km = 100 tonne-km x 0.123 kg/tkm
        self.assertEqual(self._by_step()['K3C9-1'], [('tm-02-ds', 100.0, 12.3)])

    def test_a_pcaf_investment_is_attribution_times_the_investees_emissions(self):
        # attribution = 5,000,000 / (40,000,000 + 10,000,000) = 0.10; x 1,200 tCO2e = 120 tCO2e = 120,000 kg
        self.assertEqual(self._by_step()['K3C15-1'], [('calculated-scope3-investments', 120000.0, 120000.0)])

    def test_a_leased_assets_declared_electricity_is_kwh_times_the_grid_factor(self):
        # 12,000 kWh with no facility named: the company default (Turkey distribution grid, 0.469) -> 5,628 kg.
        # A calculated entry stores its CO2e as the quantity (as the franchise and PCAF entries do).
        rows = self._by_step()['K3C8-1']
        self.assertEqual([(r[1], r[2]) for r in rows], [(5628.0, 5628.0)])
        self.assertAlmostEqual(12000 * 0.469, 5628.0, places=6)

    def test_the_dashboard_and_the_scope_split_add_up_for_these(self):
        s = self.client.get('/api/emissions/summary/?year=2025').data
        self.assertAlmostEqual(s['total_kg'], 12.3 + 120000.0 + 5628.0, places=2)
        self.assertAlmostEqual(s['scope3_kg'], s['total_kg'], places=2)                    # all three are Scope 3
        self.assertAlmostEqual(sum(c['total_co2e_kg'] for c in s['scope3_by_category']), s['scope3_kg'], places=2)
