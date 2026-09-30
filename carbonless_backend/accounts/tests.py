from django.test import TestCase
from django.contrib.auth.models import User
from rest_framework.test import APIClient


class AuthTests(TestCase):
    def test_register(self):
        client = APIClient()
        res = client.post('/api/accounts/register/', {
            'username': 'newuser', 'email': 'new@test.com',
            'password': 'StrongPass123', 'password2': 'StrongPass123',
        })
        self.assertEqual(res.status_code, 201)
        self.assertTrue(User.objects.filter(username='newuser').exists())

    def test_register_password_mismatch(self):
        client = APIClient()
        res = client.post('/api/accounts/register/', {
            'username': 'newuser2', 'email': 'new2@test.com',
            'password': 'StrongPass123', 'password2': 'WrongPass',
        })
        self.assertEqual(res.status_code, 400)

    def test_login(self):
        User.objects.create_user('loginuser', 'login@test.com', 'TestPass123')
        client = APIClient()
        res = client.post('/api/accounts/login/', {
            'username': 'loginuser', 'password': 'TestPass123',
        })
        self.assertEqual(res.status_code, 200)
        # Tokens in both cookies and body (cross-origin support)
        self.assertIn('access_token', res.cookies)
        self.assertIn('access', res.data)

    def test_profile(self):
        user = User.objects.create_user('profuser', 'prof@test.com', 'TestPass123')
        client = APIClient()
        client.force_authenticate(user=user)
        res = client.get('/api/accounts/profile/')
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.data['username'], 'profuser')

    def test_change_password(self):
        user = User.objects.create_user('pwuser', 'pw@test.com', 'OldPass123')
        client = APIClient()
        client.force_authenticate(user=user)
        res = client.post('/api/accounts/change-password/', {
            'old_password': 'OldPass123', 'new_password': 'NewPass456',
        })
        self.assertEqual(res.status_code, 200)
        user.refresh_from_db()
        self.assertTrue(user.check_password('NewPass456'))


class RegistrationCompanyTests(TestCase):
    """The company details from the signup form are saved at signup.

    The account is inactive until its email is verified, so they cannot wait
    for a post-login call — that call never ran and the details were lost.
    """

    def _register(self, username, company=None):
        body = {
            'username': username, 'email': f'{username}@test.com',
            'password': 'StrongPass123', 'password2': 'StrongPass123',
        }
        if company is not None:
            body['company'] = company
        res = APIClient().post('/api/accounts/register/', body, format='json')
        self.assertEqual(res.status_code, 201)
        user = User.objects.get(username=username)
        return res.json(), user.company_memberships.get().company

    def test_signup_company_details_are_stored(self):
        data, company = self._register('kaya', {
            'legal_entity_name': 'Kaya Tekstil A.Ş.',
            'tax_number': '1234567890',
            'country_of_headquarters': 'TR',
            'countries_of_operation': 'TR',
            'nace_code': 'C13',
            'main_activity_description': 'Kumaş üretimi',
            'number_of_employees': '51-250',
            'annual_turnover_range': '10M-50M',
            'number_of_facilities': 2,
            'has_iso_14001': True,
        })
        self.assertTrue(data['company_saved'])
        self.assertEqual(company.legal_entity_name, 'Kaya Tekstil A.Ş.')
        self.assertEqual(company.tax_number, '1234567890')
        self.assertEqual(company.nace_code, 'C13')
        self.assertTrue(company.has_iso_14001)

    def test_optional_blank_fields_fall_back_to_placeholders(self):
        # Tax number is optional on the form; a blank one must not reject the rest.
        data, company = self._register('blanktax', {
            'legal_entity_name': 'Blank Tax Ltd', 'tax_number': '',
        })
        self.assertTrue(data['company_saved'])
        self.assertEqual(company.legal_entity_name, 'Blank Tax Ltd')
        self.assertEqual(company.tax_number, '—')

    def test_without_company_details_a_placeholder_is_created(self):
        data, company = self._register('nocompany')
        self.assertFalse(data['company_saved'])
        self.assertEqual(company.legal_entity_name, "nocompany's Company")


class AuthErrorCodeTests(TestCase):
    """Auth errors carry a code so the UI can show them in Turkish."""

    def test_wrong_verification_code_reports_attempts_left(self):
        from .models import EmailVerificationToken
        user = User.objects.create_user('kod', 'kod@test.com', 'testpass123')
        token = EmailVerificationToken.objects.create(user=user)
        wrong = '000000' if token.code != '000000' else '111111'
        res = APIClient().post('/api/accounts/verify-email-code/', {'email': 'kod@test.com', 'code': wrong}, format='json')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data['code'], 'wrong_code')
        self.assertEqual(res.data['attempts_remaining'], token.MAX_ATTEMPTS - 1)

    def test_used_reset_link_has_a_code(self):
        from .models import PasswordResetToken
        user = User.objects.create_user('rst', 'rst@test.com', 'testpass123')
        token = PasswordResetToken.objects.create(user=user)
        client = APIClient()
        body = {'token': str(token.token), 'new_password': 'Yeni-Sifre-2026!'}
        self.assertEqual(client.post('/api/accounts/password-reset-confirm/', body, format='json').status_code, 200)
        res = client.post('/api/accounts/password-reset-confirm/', body, format='json')
        self.assertEqual(res.data['code'], 'link_used')


class ContactFormTests(TestCase):
    """The public contact form stores the message and emails the team."""

    URL = '/api/accounts/contact/'

    def test_message_is_stored_and_emailed_with_reply_to(self):
        from django.core import mail
        from .models import ContactMessage
        res = APIClient().post(self.URL, {
            'name': 'Ayşe', 'email': 'ayse@example.com', 'subject': 'Fiyat', 'message': 'Merhaba',
        }, format='json')
        self.assertEqual(res.status_code, 201)
        msg = ContactMessage.objects.get()
        self.assertTrue(msg.email_sent)
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].reply_to, ['ayse@example.com'])
        self.assertIn('Merhaba', mail.outbox[0].body)

    def test_missing_fields_and_bad_email_are_rejected(self):
        client = APIClient()
        res = client.post(self.URL, {'name': 'Ayşe', 'email': '', 'message': ''}, format='json')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data['code'], 'missing_fields')
        res = client.post(self.URL, {'name': 'Ayşe', 'email': 'ayse@', 'message': 'x'}, format='json')
        self.assertEqual(res.data['code'], 'invalid_email')

    def test_honeypot_is_silently_dropped(self):
        from .models import ContactMessage
        res = APIClient().post(self.URL, {
            'name': 'bot', 'email': 'bot@example.com', 'message': 'spam', 'website': 'http://spam',
        }, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertFalse(ContactMessage.objects.exists())


class SignupFacilitiesAndAutoLoginTests(TestCase):
    """Signup creates the declared facilities, and verifying the code signs the user in."""

    COMPANY = {
        'legal_entity_name': 'Yıldız Gıda Ltd.', 'tax_number': '9876543210',
        'country_of_headquarters': 'TR', 'countries_of_operation': 'TR',
        'main_activity_description': 'Unlu mamul', 'number_of_employees': '11-50',
        'annual_turnover_range': '1M - 10M ₺',
    }

    # Signup is rate-limited per IP; don't let these signups count against
    # the tests that run after this class.
    def setUp(self):
        from django.core.cache import cache
        cache.clear()

    def tearDown(self):
        from django.core.cache import cache
        cache.clear()

    def _register(self, username, facilities, language='tr'):
        res = APIClient().post('/api/accounts/register/', {
            'username': username, 'email': f'{username}@test.com',
            'password': 'StrongPass123', 'password2': 'StrongPass123',
            'company': {**self.COMPANY, 'number_of_facilities': facilities},
            'language': language,
        }, format='json')
        self.assertEqual(res.status_code, 201)
        return User.objects.get(username=username)

    def _facility_names(self, user):
        company = user.company_memberships.get().company
        return list(company.facilities.order_by('id').values_list('name', flat=True))

    def test_declared_facilities_are_created_with_numbered_names(self):
        user = self._register('zeynep', 2)
        self.assertEqual(self._facility_names(user), ['Tesis 1', 'Tesis 2'])
        self.assertEqual(user.first_name, '')

    def test_english_names_and_an_upper_bound(self):
        user = self._register('john', 500, language='en')
        names = self._facility_names(user)
        self.assertEqual(len(names), 20)
        self.assertEqual(names[0], 'Facility 1')

    def test_zero_facilities_creates_none(self):
        self.assertEqual(self._facility_names(self._register('ece', 0)), [])

    def test_correct_code_signs_the_user_in(self):
        from .models import EmailVerificationToken
        user = self._register('mert', 1)
        code = EmailVerificationToken.objects.get(user=user).code
        res = APIClient().post('/api/accounts/verify-email-code/',
                               {'email': 'mert@test.com', 'code': code}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertIn('access', res.data)
        self.assertIn('refresh', res.data)
        client = APIClient()
        client.credentials(HTTP_AUTHORIZATION=f"Bearer {res.data['access']}")
        self.assertEqual(client.get('/api/accounts/profile/').status_code, 200)

    def test_already_verified_account_gets_no_tokens(self):
        from .models import EmailVerificationToken
        user = self._register('selin', 1)
        token = EmailVerificationToken.objects.get(user=user)
        token.verify()
        res = APIClient().post('/api/accounts/verify-email-code/',
                               {'email': 'selin@test.com', 'code': '000000'}, format='json')
        self.assertEqual(res.status_code, 200)
        self.assertNotIn('access', res.data)


class DeactivatedMemberTests(TestCase):
    """A member switched off in the team keeps a login but no workspace."""

    def setUp(self):
        from rest_framework.test import APIClient
        from companies.models import Company, CompanyMembership
        self.owner = User.objects.create_user('own5', 'own5@test.com', 'pass12345')
        self.member = User.objects.create_user('mem5@test.com', 'mem5@test.com', 'pass12345',
                                               first_name='Selin', last_name='Kaya')
        self.company = Company.objects.create(
            legal_entity_name='Team5 Co', tax_number='8',
            country_of_headquarters='TR', countries_of_operation='TR',
            nace_code='', main_activity_description='x',
            number_of_employees='1-10', annual_turnover_range='x',
            number_of_facilities=1,
        )
        CompanyMembership.objects.create(user=self.owner, company=self.company, role='owner')
        self.membership = CompanyMembership.objects.create(user=self.member, company=self.company, role='data_entry')
        self.client = APIClient()

    def _profile(self, user):
        self.client.force_authenticate(user=user)
        return self.client.get('/api/accounts/profile/').data

    def test_active_member_has_access(self):
        p = self._profile(self.member)
        self.assertFalse(p['access_revoked'])
        self.assertTrue(p['permissions']['can_edit_entries'])

    def test_deactivated_member_has_no_permissions(self):
        self.membership.is_active = False
        self.membership.save()
        p = self._profile(self.member)
        self.assertTrue(p['access_revoked'])
        self.assertEqual(set(p['permissions'].values()), {False})

    def test_team_list_carries_full_name(self):
        self.client.force_authenticate(user=self.owner)
        r = self.client.get('/api/companies/memberships/')
        rows = r.data if isinstance(r.data, list) else r.data.get('results', r.data.get('members', []))
        names = {m['user_email']: m['full_name'] for m in rows}
        self.assertEqual(names['mem5@test.com'], 'Selin Kaya')


class LanguagePreferenceTests(TestCase):
    def setUp(self):
        from rest_framework.test import APIClient
        self.user = User.objects.create_user('langu', 'langu@test.com', 'pass12345')
        from .models import UserProfile
        UserProfile.objects.create(user=self.user)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_choice_is_saved_and_validated(self):
        r = self.client.patch('/api/accounts/update-profile/', {'language_preference': 'en'}, format='json')
        self.assertEqual(r.status_code, 200)
        self.user.profile.refresh_from_db()
        self.assertEqual(self.user.profile.language_preference, 'en')
        r = self.client.patch('/api/accounts/update-profile/', {'language_preference': 'de'}, format='json')
        self.assertEqual(r.status_code, 400)


class FullBackupTests(TestCase):
    def test_backup_has_company_facilities_and_inventories(self):
        from rest_framework.test import APIClient
        from companies.models import Company, CompanyMembership, Facility
        from questionnaire.models import CarbonReport, ReportStep
        user = User.objects.create_user('bak', 'bak@test.com', 'pass12345')
        company = Company.objects.create(
            legal_entity_name='Yedek A.Ş.', tax_number='9',
            country_of_headquarters='TR', countries_of_operation='TR',
            nace_code='', main_activity_description='x',
            number_of_employees='1-10', annual_turnover_range='x',
            number_of_facilities=1,
        )
        CompanyMembership.objects.create(user=user, company=company, role='owner')
        Facility.objects.create(company=company, name='Gebze Fabrika')
        report = CarbonReport.objects.create(company=company, created_by=user, reporting_year=2024, title='Env 2024')
        ReportStep.objects.create(report=report, step_id='4A-1', answer={'answer': {'Merkez': '1000 kWh'}})
        c = APIClient()
        c.force_authenticate(user=user)
        d = c.get('/api/emissions/export-all/').data
        self.assertEqual(d['company']['legal_entity_name'], 'Yedek A.Ş.')
        self.assertEqual([f['name'] for f in d['facilities']], ['Gebze Fabrika'])
        self.assertEqual(d['inventories'][0]['answers']['4A-1'], {'answer': {'Merkez': '1000 kWh'}})
        r = c.get('/api/emissions/export-all/?file=xlsx&lang=tr')
        self.assertEqual(r.status_code, 200)
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(b''.join(r.streaming_content) if r.streaming else r.content))
        self.assertEqual(wb.sheetnames, ['Şirket', 'Tesisler', 'Emisyon Kayıtları', 'Hedefler',
                                         'Envanterler', 'Anket Cevapları'])
        self.assertEqual(wb['Anket Cevapları']['C2'].value, '4A-1')


class RateLimitKeyTests(TestCase):
    """Every visitor arrives through the same proxy address, so limits must be
    counted per account (sign-in) or per forwarded client address — never on
    the proxy's own address, which would lock the whole site at once."""

    def setUp(self):
        from types import SimpleNamespace
        from unittest import mock
        from django.core.cache import cache
        cache.clear()
        self.user = User.objects.create_user('real', 'real@test.com', 'StrongPass123')
        # django-ratelimit counts in fixed clock windows; attempts that straddle
        # a window boundary start a new count and the limit is never reached.
        # Freeze the clock it reads (only there) so these tests are exact.
        patcher = mock.patch('django_ratelimit.core.time', SimpleNamespace(time=lambda: 1_800_000_000.0))
        patcher.start()
        self.addCleanup(patcher.stop)

    def tearDown(self):
        from django.core.cache import cache
        cache.clear()

    def _login(self, username, password, **extra):
        return APIClient().post('/api/accounts/login/', {'username': username, 'password': password},
                                format='json', **extra)

    def test_others_failed_logins_do_not_lock_a_user_out(self):
        for _ in range(12):
            self._login('attacker-target', 'wrong')
        res = self._login('real', 'StrongPass123')
        self.assertEqual(res.status_code, 200)

    def test_too_many_failed_logins_answer_429_with_code(self):
        for _ in range(10):
            self._login('real', 'wrong')
        res = self._login('REAL ', 'StrongPass123')
        self.assertEqual(res.status_code, 429)
        self.assertEqual(res.json()['code'], 'rate_limited')

    def test_forwarded_client_address_is_used_behind_the_proxy(self):
        from django.test import override_settings
        with override_settings(RATELIMIT_TRUSTED_PROXIES=1):
            for _ in range(10):
                res = APIClient().post('/api/accounts/password-reset/', {'email': 'x@test.com'},
                                       format='json', HTTP_X_FORWARDED_FOR='203.0.113.5')
                self.assertEqual(res.status_code, 200)
            blocked = APIClient().post('/api/accounts/password-reset/', {'email': 'x@test.com'},
                                       format='json', HTTP_X_FORWARDED_FOR='203.0.113.5')
            self.assertEqual(blocked.status_code, 429)
            other = APIClient().post('/api/accounts/password-reset/', {'email': 'x@test.com'},
                                     format='json', HTTP_X_FORWARDED_FOR='198.51.100.7')
            self.assertEqual(other.status_code, 200)

    def test_spoofed_leftmost_forwarded_entry_is_ignored(self):
        from carbonless_api.client_ip import client_ip
        from django.test import RequestFactory, override_settings
        req = RequestFactory().get('/', HTTP_X_FORWARDED_FOR='1.2.3.4, 203.0.113.5', REMOTE_ADDR='10.0.0.1')
        with override_settings(RATELIMIT_TRUSTED_PROXIES=1):
            self.assertEqual(client_ip(req), '203.0.113.5')
        with override_settings(RATELIMIT_TRUSTED_PROXIES=0):
            self.assertEqual(client_ip(req), '10.0.0.1')


class AccountEmailLanguageTests(TestCase):
    """The verification-code and password-reset e-mails come in the user's
    language: the page's language when the request sends one, else the
    account's preference. They used to be English only."""

    def setUp(self):
        from django.core.cache import cache
        cache.clear()

    def _user(self, lang):
        from .models import UserProfile
        user = User.objects.create_user(f'u{lang}', f'u{lang}@test.com', 'StrongPass123', first_name='Ayşe')
        UserProfile.objects.create(user=user, language_preference=lang)
        return user

    def test_password_reset_follows_preference_then_page_language(self):
        from django.core import mail
        self._user('tr')
        APIClient().post('/api/accounts/password-reset/', {'email': 'utr@test.com'}, format='json')
        self.assertEqual(mail.outbox[-1].subject, 'Carbonless şifrenizi sıfırlayın')
        self.assertIn('Merhaba Ayşe', mail.outbox[-1].body)
        self.assertIn('/reset-password?token=', mail.outbox[-1].body)
        APIClient().post('/api/accounts/password-reset/', {'email': 'utr@test.com', 'language': 'en'}, format='json')
        self.assertEqual(mail.outbox[-1].subject, 'Reset your Carbonless password')
        self.assertIn('Hi Ayşe', mail.outbox[-1].body)

    def test_signup_code_uses_the_signup_language(self):
        from django.core import mail
        body = {'username': 'mehmet', 'email': 'mehmet@test.com',
                'password': 'StrongPass123', 'password2': 'StrongPass123', 'language': 'tr'}
        self.assertEqual(APIClient().post('/api/accounts/register/', body, format='json').status_code, 201)
        self.assertTrue(mail.outbox[-1].subject.startswith('Carbonless doğrulama kodunuz: '))
        self.assertIn('Doğrulama kodunuz', mail.outbox[-1].body)

    def test_resend_code_uses_the_page_language(self):
        from django.core import mail
        user = self._user('tr')
        user.is_active = False
        user.save()
        APIClient().post('/api/accounts/resend-verification/', {'email': 'utr@test.com', 'language': 'en'}, format='json')
        self.assertIn('is your Carbonless verification code', mail.outbox[-1].subject)
        APIClient().post('/api/accounts/resend-verification/', {'email': 'utr@test.com'}, format='json')
        self.assertTrue(mail.outbox[-1].subject.startswith('Carbonless doğrulama kodunuz: '))


class NotificationCompanyTests(TestCase):
    """A member of several companies sees a company's notices only while
    working in it; company-less notices show everywhere."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        from .models import Notification, UserProfile
        kw = dict(country_of_headquarters='TR', countries_of_operation='TR', main_activity_description='x',
                  number_of_employees='1-10', annual_turnover_range='x', number_of_facilities=1)
        self.a = Company.objects.create(legal_entity_name='A A.Ş.', tax_number='1111111111', **kw)
        self.b = Company.objects.create(legal_entity_name='B A.Ş.', tax_number='2222222222', **kw)
        self.user = User.objects.create_user('multi', 'multi@test.com', 'testpass123')
        self.profile = UserProfile.objects.create(user=self.user, active_company=self.a)
        for c in (self.a, self.b):
            CompanyMembership.objects.create(company=c, user=self.user, role='owner')
        for company, title in ((self.a, 'about A'), (self.b, 'about B'), (None, 'billing')):
            Notification.objects.create(user=self.user, company=company, notification_type='system',
                                        title=title, message='x')
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def _titles(self):
        return sorted(n['title'] for n in self.client.get('/api/accounts/notifications/').data)

    def test_list_count_and_mark_all_follow_the_active_company(self):
        self.assertEqual(self._titles(), ['about A', 'billing'])
        self.assertEqual(self.client.get('/api/accounts/notifications/unread-count/').data['unread_count'], 2)
        self.client.post('/api/accounts/notifications/read/', {}, format='json')
        from .models import Notification
        self.assertFalse(Notification.objects.get(title='about B').is_read)
        self.profile.active_company = self.b
        self.profile.save()
        self.assertEqual(self._titles(), ['about B', 'billing'])
        self.assertEqual(self.client.get('/api/accounts/notifications/unread-count/').data['unread_count'], 1)

    def test_entry_notice_is_tagged_with_the_entrys_company(self):
        from emissions.models import EmissionEntry, EmissionFactor
        from emissions.notifications import notify_entry_submitted
        from companies.models import CompanyMembership
        from .models import Notification
        member = User.objects.create_user('de', 'de@test.com', 'testpass123')
        CompanyMembership.objects.create(company=self.b, user=member, role='data_entry')
        f = EmissionFactor.objects.create(slug='n-f', name='F', scope='scope1', category='stationary_combustion',
                                          country='global', unit='kg', factor_kg_co2e=1, source='generic')
        entry = EmissionEntry.objects.create(user=member, company=self.b, emission_factor=f, year=2026, month=1,
                                             quantity=1, calculated_co2e_kg=1, status='submitted')
        notify_entry_submitted(entry)
        self.assertEqual(Notification.objects.get(notification_type='entry_submitted').company, self.b)


class HistoryTextTests(TestCase):
    def test_turkish_reader_gets_turkish_name_month_unit_and_number(self):
        from .history_text import localize_detail
        names = {'Natural Gas (Turkey)': 'Doğal Gaz (Türkiye)'}
        self.assertEqual(
            localize_detail('Natural Gas (Turkey) · 2026/09 · 1250.5 gj · Kanıt eklendi: f.pdf', 'tr', names),
            'Doğal Gaz (Türkiye) · Eylül 2026 · 1.250,5 GJ · Kanıt eklendi: f.pdf')

    def test_english_reader_and_changes(self):
        from .history_text import localize_detail
        self.assertEqual(
            localize_detail('Grid · 2026/08 · 45 kwh → Grid · 2026/08 · 50 kwh · Kanıt kaldırıldı: a.pdf', 'en', {}),
            'Grid · August 2026 · 45 kWh → Grid · August 2026 · 50 kWh · Proof removed: a.pdf')

    def test_legacy_and_unknown_text(self):
        from .history_text import localize_detail
        self.assertEqual(localize_detail('Updated emission entry: Grid (12 kwh)', 'tr', {'Grid': 'Şebeke'}),
                         'Şebeke · 12 kWh')
        self.assertEqual(localize_detail('something else', 'tr', {}), 'something else')
