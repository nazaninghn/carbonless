from decimal import Decimal

from django.core.exceptions import ValidationError
from django.test import TestCase
from django.contrib.auth.models import User
from rest_framework.test import APIClient
from .gas_split_data import fuel_gas_shares
from .models import EmissionFactor, EmissionEntry


class EmissionFactorTests(TestCase):
    def setUp(self):
        self.factor = EmissionFactor.objects.create(
            slug='test-gas', name='Test Gas', name_tr='Test Gaz',
            scope='scope1', category='stationary_combustion', country='global',
            unit='kg', factor_kg_co2e=2.5, year=2024, source='generic',
            is_active=True, is_default=True,
        )

    def test_factor_created(self):
        self.assertEqual(EmissionFactor.objects.count(), 1)
        self.assertEqual(self.factor.factor_kg_co2e, 2.5)

    def test_factor_str(self):
        self.assertIn('Test Gas', str(self.factor))


class GasSplitTests(TestCase):
    """The per-gas breakdown ISO 14064-1 reporting needs."""

    def _factor(self, **kw):
        defaults = dict(
            slug='split-test', name='Split Test', scope='scope1',
            category='stationary_combustion', country='global', unit='kg',
            factor_kg_co2e=Decimal('100'), year=2024, source='generic',
        )
        defaults.update(kw)
        return EmissionFactor(**defaults)

    def test_no_split_reads_as_none(self):
        """A factor with no breakdown must not look like an all-zero one."""
        f = self._factor()
        self.assertIsNone(f.gas_split())
        self.assertIsNone(f.gas_split_shares())

    def test_shares_sum_to_one(self):
        f = self._factor(
            factor_co2_kg_co2e=Decimal('97'), factor_ch4_kg_co2e=Decimal('1'),
            factor_n2o_kg_co2e=Decimal('2'), gas_split_basis='apportioned')
        shares = f.gas_split_shares()
        self.assertEqual(sorted(shares), ['CH4', 'CO2', 'N2O'])
        self.assertAlmostEqual(float(sum(shares.values())), 1.0, places=9)
        self.assertAlmostEqual(float(shares['CO2']), 0.97, places=9)

    def test_gases_with_no_contribution_are_omitted(self):
        f = self._factor(factor_sf6_kg_co2e=Decimal('100'),
                         gas_split_basis='single_gas')
        self.assertEqual(f.gas_split(), {'SF6': Decimal('100')})

    def test_split_must_add_up_to_the_factor(self):
        f = self._factor(factor_co2_kg_co2e=Decimal('50'),
                         gas_split_basis='apportioned')
        with self.assertRaises(ValidationError) as ctx:
            f.clean()
        self.assertIn('factor_kg_co2e', ctx.exception.message_dict)

    def test_rounding_drift_is_tolerated(self):
        f = self._factor(factor_co2_kg_co2e=Decimal('99.995'),
                         gas_split_basis='apportioned')
        f.clean()  # within tolerance — must not raise

    def test_split_requires_a_stated_basis(self):
        """A breakdown with no provenance cannot go in an assurance document."""
        f = self._factor(factor_co2_kg_co2e=Decimal('100'))
        with self.assertRaises(ValidationError) as ctx:
            f.clean()
        self.assertIn('gas_split_basis', ctx.exception.message_dict)

    def test_seeded_combustion_factor_keeps_its_total(self):
        """Apportioning must divide a factor, never change it."""
        shares = fuel_gas_shares('diesel_onroad')
        self.assertAlmostEqual(sum(shares.values()), 1.0, places=9)
        # CO2 dominates, but CH4 and N2O are both present and non-zero.
        self.assertGreater(shares['CO2'], 0.9)
        self.assertGreater(shares['CH4'], 0)
        self.assertGreater(shares['N2O'], 0)


class EmissionEntryTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('testuser', 'test@test.com', 'testpass123')
        self.factor = EmissionFactor.objects.create(
            slug='test-elec', name='Test Electricity', name_tr='Test Elektrik',
            scope='scope2', category='electricity', country='turkey',
            unit='kwh', factor_kg_co2e=0.4199, year=2023, source='atom_kablo',
            is_active=True, is_default=True,
        )

    def test_entry_auto_calculate(self):
        entry = EmissionEntry.objects.create(
            user=self.user, emission_factor=self.factor,
            year=2026, month=1, quantity=1000,
        )
        self.assertAlmostEqual(float(entry.calculated_co2e_kg), 419.9, places=1)

    def test_entry_tonne(self):
        entry = EmissionEntry.objects.create(
            user=self.user, emission_factor=self.factor,
            year=2026, month=1, quantity=1000,
        )
        self.assertAlmostEqual(float(entry.calculated_co2e_tonne), 0.4199, places=3)


class CalculatorTests(TestCase):
    def setUp(self):
        self.factor = EmissionFactor.objects.create(
            slug='test-diesel', name='Test Diesel', name_tr='Test Dizel',
            scope='scope1', category='mobile_combustion', country='turkey',
            unit='liters', factor_kg_co2e=2.696, year=2023, source='atom_kablo',
            is_active=True, is_default=True,
        )

    def test_calculate_by_id(self):
        from .calculator import calculate_emissions
        result = calculate_emissions(self.factor.pk, 100)
        self.assertAlmostEqual(result['emissions_kg'], 269.6, places=1)
        self.assertEqual(result['scope'], 'scope1')

    def test_calculate_negative(self):
        from .calculator import calculate_emissions
        result = calculate_emissions(self.factor.pk, -10)
        self.assertIn('error', result)

    def test_get_factor_lookup(self):
        from .calculator import get_emission_factor
        f = get_emission_factor('test-diesel', 'turkey', 'mobile_combustion')
        self.assertIsNotNone(f)
        self.assertEqual(float(f.factor_kg_co2e), 2.696)

    def test_get_factor_not_found(self):
        from .calculator import get_emission_factor
        f = get_emission_factor('nonexistent', 'turkey', 'electricity')
        self.assertIsNone(f)


class APITests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user('apiuser', 'api@test.com', 'testpass123')
        from companies.models import Company, CompanyMembership
        self.company = Company.objects.create(
            legal_entity_name='Test Co', tax_number='123',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='Test', number_of_employees='1-10',
            annual_turnover_range='<1M', number_of_facilities=1,
        )
        CompanyMembership.objects.create(company=self.company, user=self.user, role='owner')
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.factor = EmissionFactor.objects.create(
            slug='api-test', name='API Test Factor', name_tr='API Test',
            scope='scope2', category='electricity', country='turkey',
            unit='kwh', factor_kg_co2e=0.4199, year=2023, source='atom_kablo',
            is_active=True, is_default=True,
        )

    def test_list_factors(self):
        res = self.client.get('/api/emissions/factors/')
        self.assertEqual(res.status_code, 200)

    def test_create_entry(self):
        res = self.client.post('/api/emissions/entries/', {
            'emission_factor': self.factor.pk,
            'year': 2026, 'month': 1, 'quantity': 500,
        })
        self.assertEqual(res.status_code, 201)

    def test_summary(self):
        EmissionEntry.objects.create(
            user=self.user, company=self.company, emission_factor=self.factor,
            year=2026, month=1, quantity=1000, status='approved',
        )
        res = self.client.get('/api/emissions/summary/?year=2026')
        self.assertEqual(res.status_code, 200)
        self.assertGreater(res.data['total_kg'], 0)

    def test_calculate_endpoint(self):
        res = self.client.post('/api/emissions/calculate/', {
            'factor_id': self.factor.pk, 'activity_data': 1000,
        })
        self.assertEqual(res.status_code, 200)
        self.assertAlmostEqual(res.data['emissions_kg'], 419.9, places=1)


class ProofAndAccountDeletionTests(TestCase):
    """Proof uploads are stored, and company data outlives the member who entered it."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        from accounts.models import UserProfile
        self.factor = EmissionFactor.objects.create(
            slug='test-gas-proof', name='Test Gas', name_tr='Test Gaz',
            scope='scope1', category='stationary_combustion', country='global',
            unit='kg', factor_kg_co2e=2.5, year=2024, source='generic',
            is_active=True, is_default=True,
        )
        self.company = Company.objects.create(
            legal_entity_name='Kaya Tekstil A.Ş.', tax_number='1234567890',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        self.users = {}
        for name, role in [('aylin', 'owner'), ('ali', 'data_entry')]:
            u = User.objects.create_user(name, f'{name}@test.com', 'testpass123')
            UserProfile.objects.create(user=u, active_company=self.company)
            CompanyMembership.objects.create(company=self.company, user=u, role=role)
            self.users[name] = u
        self.client = APIClient()

    def _as(self, name):
        self.client.force_authenticate(user=self.users[name])

    def _entry(self, **extra):
        return self.client.post('/api/emissions/entries/', {
            'emission_factor': self.factor.id, 'year': 2026, 'month': 3, 'quantity': 10, **extra,
        }, format='multipart')

    def _file(self, name, size=100):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, b'%PDF' + b'0' * size, content_type='application/pdf')

    def test_proof_document_is_saved_and_downloadable(self):
        import shutil, tempfile
        media = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, media, True)
        with self.settings(MEDIA_ROOT=media):
            self._as('aylin')
            res = self._entry(proof_document=self._file('fatura.pdf'))
            self.assertEqual(res.status_code, 201, res.data)
            self.assertTrue(res.data['proof_document'])
            dl = self.client.get(f"/api/emissions/entries/{res.data['id']}/proof/")
            self.assertEqual(dl.status_code, 200)
            self.assertTrue(b''.join(dl.streaming_content).startswith(b'%PDF'))

    def test_bad_proof_document_is_a_400_not_a_500(self):
        self._as('aylin')
        self.assertEqual(self._entry(proof_document=self._file('virus.exe')).status_code, 400)
        self.assertEqual(self._entry(proof_document=self._file('big.pdf', 10 * 1024 * 1024 + 1)).status_code, 400)
        self.assertEqual(EmissionEntry.objects.count(), 0)

    def _delete_account(self, name):
        self._as(name)
        return self.client.delete('/api/accounts/delete-account/', {'password': 'testpass123'}, format='json')

    def test_member_leaving_keeps_their_entries_in_the_company(self):
        self._as('ali')
        entry_id = self._entry().data['id']
        self.assertEqual(self._delete_account('ali').status_code, 200)
        entry = EmissionEntry.objects.get(id=entry_id)
        self.assertEqual(entry.company, self.company)
        self.assertIsNone(entry.user)
        self._as('aylin')
        listed = self.client.get('/api/emissions/entries/').data
        rows = listed['results'] if isinstance(listed, dict) else listed
        self.assertIn(entry_id, [r['id'] for r in rows])

    def test_owner_needs_an_admin_or_manager_to_take_over(self):
        res = self._delete_account('aylin')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data['code'], 'owner_successor_needed')
        self.assertTrue(User.objects.filter(username='aylin').exists())

    def test_owner_leaving_hands_the_company_to_a_manager(self):
        from companies.models import CompanyMembership
        CompanyMembership.objects.filter(user=self.users['ali']).update(role='manager')
        self._as('aylin')
        entry_id = self._entry().data['id']
        self.assertEqual(self._delete_account('aylin').status_code, 200)
        self.assertEqual(CompanyMembership.objects.get(user=self.users['ali']).role, 'owner')
        self.assertTrue(EmissionEntry.objects.filter(id=entry_id, company=self.company).exists())

    def test_sole_owner_leaving_deletes_their_company(self):
        from companies.models import Company, CompanyMembership
        CompanyMembership.objects.filter(user=self.users['ali']).delete()
        self._as('aylin')
        self._entry()
        self.assertEqual(self._delete_account('aylin').status_code, 200)
        self.assertFalse(Company.objects.filter(id=self.company.id).exists())
        self.assertEqual(EmissionEntry.objects.count(), 0)

    def test_wrong_password_has_a_code(self):
        self._as('aylin')
        res = self.client.delete('/api/accounts/delete-account/', {'password': 'nope'}, format='json')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data['code'], 'wrong_password')

    def test_form_entry_status_follows_role_like_chat(self):
        self._as('aylin')
        self.assertEqual(self._entry().data['status'], 'approved')
        self._as('ali')
        self.assertEqual(self._entry(quantity=11).data['status'], 'submitted')


class ApprovalNotificationAndDuplicateTests(TestCase):
    """Approvers hear about new entries, authors hear back, and duplicates are flagged."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        from accounts.models import UserProfile
        self.factor = EmissionFactor.objects.create(
            slug='test-gas-dup', name='Test Gas', name_tr='Test Gaz',
            scope='scope1', category='stationary_combustion', country='global',
            unit='kg', factor_kg_co2e=2.5, year=2024, source='generic',
            is_active=True, is_default=True,
        )
        self.company = Company.objects.create(
            legal_entity_name='Kaya Tekstil A.Ş.', tax_number='1234567890',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        self.users = {}
        for name, role, lang in [('aylin', 'owner', 'tr'), ('ali', 'data_entry', 'en')]:
            u = User.objects.create_user(name, f'{name}@test.com', 'testpass123')
            UserProfile.objects.create(user=u, active_company=self.company, language_preference=lang)
            CompanyMembership.objects.create(company=self.company, user=u, role=role)
            self.users[name] = u
        self.client = APIClient()

    def _as(self, name):
        self.client.force_authenticate(user=self.users[name])

    def _entry(self, **extra):
        return self.client.post('/api/emissions/entries/', {
            'emission_factor': self.factor.id, 'year': 2026, 'month': 3, 'quantity': 10, **extra,
        }, format='json')

    def _notes(self, name, kind):
        from accounts.models import Notification
        return list(Notification.objects.filter(user=self.users[name], notification_type=kind))

    def test_submitted_entry_notifies_the_approver_in_their_language(self):
        self._as('ali')
        self.assertEqual(self._entry().status_code, 201)
        notes = self._notes('aylin', 'entry_submitted')
        self.assertEqual(len(notes), 1)
        self.assertEqual(notes[0].title, 'Onay bekleyen kayıt')
        self.assertIn('Test Gaz', notes[0].message)
        # No name on the account: the e-mail stands in for it, as in the
        # Onay Bekleyenler list ("Giren: …").
        self.assertTrue(notes[0].message.startswith('ali@test.com '))

    def test_notice_names_the_author_by_full_name(self):
        self.users['ali'].first_name, self.users['ali'].last_name = 'Ali', 'Kaya'
        self.users['ali'].save()
        self._as('ali')
        self._entry()
        self.assertTrue(self._notes('aylin', 'entry_submitted')[0].message.startswith('Ali Kaya yeni bir kayıt ekledi'))

    def test_owner_entry_needs_no_approval_notice(self):
        self._as('aylin')
        self._entry()
        self.assertEqual(self._notes('aylin', 'entry_submitted'), [])

    def test_rejection_tells_the_author_why_and_editing_resubmits(self):
        self._as('ali')
        entry_id = self._entry().data['id']
        self._as('aylin')
        res = self.client.post(f'/api/emissions/entries/{entry_id}/approve/',
                               {'action': 'reject', 'reason': 'Invoice missing'}, format='json')
        self.assertEqual(res.status_code, 200)
        notes = self._notes('ali', 'entry_rejected')
        self.assertEqual(len(notes), 1)
        self.assertIn('Invoice missing', notes[0].message)

        self._as('ali')
        res = self.client.patch(f'/api/emissions/entries/{entry_id}/', {'quantity': 12}, format='json')
        self.assertEqual(res.status_code, 200)
        entry = EmissionEntry.objects.get(id=entry_id)
        self.assertEqual(entry.status, 'submitted')
        self.assertEqual(entry.rejected_reason, '')
        self.assertEqual(len(self._notes('aylin', 'entry_submitted')), 2)

    def test_approval_tells_the_author(self):
        self._as('ali')
        entry_id = self._entry().data['id']
        self._as('aylin')
        self.client.post(f'/api/emissions/entries/{entry_id}/approve/', {'action': 'approve'}, format='json')
        self.assertEqual(len(self._notes('ali', 'entry_approved')), 1)

    def test_switch_off_stops_emails_but_not_in_app_notices(self):
        # The switch is for e-mail: turned off, a rejection still reaches the
        # author in the app (they never learned of it before).
        from django.core import mail
        for name in ('aylin', 'ali'):
            self.users[name].profile.notify_approvals = False
            self.users[name].profile.save()
        self._as('ali')
        entry_id = self._entry().data['id']
        self.assertEqual(len(self._notes('aylin', 'entry_submitted')), 1)
        self._as('aylin')
        self.client.post(f'/api/emissions/entries/{entry_id}/approve/',
                         {'action': 'reject', 'reason': 'Fatura eksik'}, format='json')
        [note] = self._notes('ali', 'entry_rejected')
        self.assertIn('Fatura eksik', note.message)
        self.assertEqual(mail.outbox, [])

    def test_same_entry_twice_is_flagged_unless_confirmed(self):
        self._as('aylin')
        self.assertEqual(self._entry().status_code, 201)
        res = self._entry()
        self.assertEqual(res.status_code, 409)
        self.assertEqual(res.data['code'], 'possible_duplicate')
        self.assertEqual(EmissionEntry.objects.count(), 1)
        self.assertEqual(self._entry(confirm_duplicate=True).status_code, 201)
        self.assertEqual(EmissionEntry.objects.count(), 2)

    def test_different_month_or_rejected_entry_is_not_a_duplicate(self):
        self._as('aylin')
        self._entry()
        self.assertEqual(self._entry(month=4).status_code, 201)
        EmissionEntry.objects.filter(month=3).update(status='draft')
        self.assertEqual(self._entry().status_code, 201)


class YearsAndFacilityTests(TestCase):
    """The summary lists the years that have data; deleting a facility keeps its entries."""

    def setUp(self):
        from companies.models import Company, CompanyMembership, Facility
        from accounts.models import UserProfile
        self.factor = EmissionFactor.objects.create(
            slug='test-gas-years', name='Test Gas', name_tr='Test Gaz',
            scope='scope1', category='stationary_combustion', country='global',
            unit='kg', factor_kg_co2e=2.5, year=2024, source='generic',
            is_active=True, is_default=True,
        )
        self.company = Company.objects.create(
            legal_entity_name='Kaya Tekstil A.Ş.', tax_number='1234567890',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        self.user = User.objects.create_user('aylin', 'aylin@test.com', 'testpass123')
        UserProfile.objects.create(user=self.user, active_company=self.company)
        CompanyMembership.objects.create(company=self.company, user=self.user, role='owner')
        self.facility = Facility.objects.create(company=self.company, name='Fabrika 1')
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def _entry(self, year, **extra):
        return self.client.post('/api/emissions/entries/', {
            'emission_factor': self.factor.id, 'year': year, 'month': 3, 'quantity': 10, **extra,
        }, format='json')

    def test_summary_lists_years_with_data_newest_first(self):
        self._entry(2024)
        self._entry(2026)
        data = self.client.get('/api/emissions/summary/?year=2025').data
        self.assertEqual(data['years_with_data'], [2026, 2024])
        self.assertEqual(data['total_kg'], 0)

    def test_facility_shows_entry_count_and_delete_keeps_entries(self):
        entry_id = self._entry(2026, facility=self.facility.id).data['id']
        listed = self.client.get('/api/companies/facilities/').data
        rows = listed['results'] if isinstance(listed, dict) else listed
        self.assertEqual(rows[0]['entry_count'], 1)
        res = self.client.patch(f'/api/companies/facilities/{self.facility.id}/', {'name': 'Fabrika A'}, format='json')
        self.assertEqual(res.data['name'], 'Fabrika A')
        self.assertEqual(self.client.delete(f'/api/companies/facilities/{self.facility.id}/').status_code, 204)
        entry = EmissionEntry.objects.get(id=entry_id)
        self.assertIsNone(entry.facility)


class AnnualQuestionnaireEntryTests(TestCase):
    """Questionnaire entries are annual amounts stored with month=1: they are
    kept out of the monthly breakdown (and its January) and shown as annual."""

    _entry = YearsAndFacilityTests._entry

    def setUp(self):
        YearsAndFacilityTests.setUp(self)
        self._entry(2025)  # a monthly entry: March, 25 kg
        EmissionEntry.objects.create(
            user=self.user, company=self.company, emission_factor=self.factor, year=2025, month=1,
            quantity=100, status='approved', description='Questionnaire step 3A-5 · Merkez')

    def test_summary_keeps_annual_out_of_months(self):
        data = self.client.get('/api/emissions/summary/?year=2025').data
        self.assertEqual(data['total_kg'], 275)
        self.assertEqual(data['annual_kg'], 250)
        self.assertEqual(data['monthly'][0]['total_kg'], 0)    # January
        self.assertEqual(data['monthly'][2]['total_kg'], 25)   # March

    def test_summary_splits_by_facility(self):
        # (same source as the annual questionnaire entry: saved knowingly)
        r = self._entry(2025, facility=self.facility.id, quantity=4, confirm_duplicate=True)   # 10 kg at Fabrika 1
        self.assertEqual(r.status_code, 201, r.content)
        data = self.client.get('/api/emissions/summary/?year=2025').data
        self.assertEqual(data['by_facility'], [{'facility_id': self.facility.id, 'facility_name': 'Fabrika 1',
                                                'total_kg': 10.0}])
        self.assertEqual(data['unassigned_kg'], 275.0)

    def test_excel_month_column_says_annual(self):
        import io
        from openpyxl import load_workbook
        res = self.client.get('/api/emissions/export-excel/?year=2025&lang=tr')
        rows = list(load_workbook(io.BytesIO(res.content)).active.iter_rows(values_only=True))
        self.assertEqual(sorted(r[3] for r in rows[1:]), ['Mart', 'Yıllık'])

    def test_chat_month_answer_leaves_annual_out(self):
        from chat.local_answers import _month_answer
        text = _month_answer(self.company, 2025, 1, 'tr')
        self.assertIn('aylık emisyon verisi yok', text)
        self.assertIn('yıllık envanter (0,25 tCO₂e)', text)

    def test_profile_follows_selected_year_with_sector(self):
        from questionnaire.models import CarbonReport, ReportStep
        self.company.nace_code = 'C20'
        self.company.save()
        older = CarbonReport.objects.create(company=self.company, created_by=self.user, title='Envanter 2025',
                                            reporting_year=2025, status='completed')
        CarbonReport.objects.create(company=self.company, created_by=self.user, title='Envanter 2024',
                                    reporting_year=2024, status='completed')
        ReportStep.objects.create(report=older, step_id='B3', answer={'employee_band': '51-250'})
        profile = self.client.get('/api/emissions/summary/?year=2025').data['questionnaire_profile']
        self.assertEqual((profile['title'], profile['nace_code'], profile['employee_band']),
                         ('Envanter 2025', 'C20', '51-250'))
        # A year without an inventory: the newest one, saying which it is.
        profile = self.client.get('/api/emissions/summary/?year=2026').data['questionnaire_profile']
        self.assertEqual(profile['title'], 'Envanter 2024')

    def test_summary_says_this_years_inventory_is_in_progress(self):
        from questionnaire.models import CarbonReport
        draft = CarbonReport.objects.create(company=self.company, created_by=self.user, title='Envanter 2023',
                                            reporting_year=2023, status='in_progress',
                                            client_progress={'answered': 20, 'total': 50})
        info = self.client.get('/api/emissions/summary/?year=2023').data['inventory_in_progress']
        self.assertEqual((info['report_id'], info['year'], info['percent']), (draft.id, 2023, 40))
        # Another year, or a year whose inventory is finished: nothing to say.
        self.assertIsNone(self.client.get('/api/emissions/summary/?year=2025').data['inventory_in_progress'])
        draft.status = 'completed'
        draft.save()
        self.assertIsNone(self.client.get('/api/emissions/summary/?year=2023').data['inventory_in_progress'])


class ExcelExportLanguageTests(TestCase):
    """The Excel export follows the UI language and shows each entry's status."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        from accounts.models import UserProfile
        factor = EmissionFactor.objects.create(
            slug='test-gas-xlsx', name='Natural Gas (Turkey)', name_tr='Doğal Gaz (Türkiye)',
            scope='scope1', category='stationary_combustion', country='global',
            unit='gj', factor_kg_co2e=56.211, year=2024, source='generic',
            is_active=True, is_default=True,
        )
        company = Company.objects.create(
            legal_entity_name='Kaya Tekstil A.Ş.', tax_number='1234567890',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        self.user = User.objects.create_user('xlsx', 'xlsx@test.com', 'testpass123')
        UserProfile.objects.create(user=self.user, active_company=company)
        CompanyMembership.objects.create(company=company, user=self.user, role='owner')
        EmissionEntry.objects.create(user=self.user, company=company, emission_factor=factor,
                                     year=2026, month=3, quantity=10, status='draft',
                                     rejected_reason='Fatura eksik')
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def _rows(self, lang):
        import io
        from openpyxl import load_workbook
        res = self.client.get(f'/api/emissions/export-excel/?year=2026&lang={lang}')
        self.assertEqual(res.status_code, 200)
        ws = load_workbook(io.BytesIO(res.content)).active
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def test_turkish_export(self):
        header, row = self._rows('tr')
        self.assertEqual(header[0], 'Kaynak')
        self.assertEqual(row[:4], ['Doğal Gaz (Türkiye)', 'Kapsam 1', 'Sabit Yanma', 'Mart'])
        self.assertEqual(row[-2:], ['Reddedildi', 'Fatura eksik'])

    def test_unit_and_system_description_are_readable(self):
        EmissionEntry.objects.update(description='AI Chat: natural_gas 10 gj')
        header, row = self._rows('tr')
        self.assertEqual(row[header.index('Birim')], 'GJ')
        self.assertEqual(row[header.index('Açıklama')], 'AI sohbetinden')
        EmissionEntry.objects.update(description='Questionnaire step 3A-5 · Merkez')
        header, row = self._rows('en')
        self.assertEqual(row[header.index('Description')], 'From the questionnaire — Question 44 · Merkez')

    def test_english_export(self):
        header, row = self._rows('en')
        self.assertEqual(header[-2:], ['Status', 'Rejection reason'])
        self.assertEqual(row[1:3], ['Scope 1', 'Stationary Combustion'])
        self.assertEqual(row[-2], 'Rejected')

    def _csv(self, lang):
        res = self.client.get(f'/api/emissions/export-csv/?year=2026&lang={lang}')
        self.assertEqual(res.status_code, 200)
        text = res.content.decode('utf-8')
        self.assertTrue(text.startswith('\ufeff'))
        return [line for line in text.lstrip('\ufeff').splitlines() if line]

    def test_turkish_csv_opens_in_turkish_excel(self):
        header, row = self._csv('tr')
        self.assertTrue(header.startswith('Kaynak;Kapsam;Kategori;Ay;Miktar'))
        cells = row.split(';')
        self.assertEqual(cells[:5], ['Doğal Gaz (Türkiye)', 'Kapsam 1', 'Sabit Yanma', 'Mart', '10'])
        self.assertEqual(cells[6], '56,211')       # decimal comma
        self.assertEqual(cells[8], '0,56211')      # no float noise
        self.assertEqual(cells[-2:], ['Reddedildi', 'Fatura eksik'])

    def test_english_csv(self):
        header, row = self._csv('en')
        self.assertTrue(header.startswith('Source,Scope,Category'))
        self.assertIn('Rejected', row)
        self.assertIn('56.211', row)


class TargetAndCustomRequestRoleTests(TestCase):
    """Targets are the company's commitment: only owner/admin/manager change
    them. A custom request is withdrawn or edited only while pending, by its
    sender or an owner/admin/manager."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        from accounts.models import UserProfile
        self.company = Company.objects.create(
            legal_entity_name='Rol A.Ş.', tax_number='1234567891',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        self.users = {}
        for name, role in [('sahip', 'owner'), ('mudur', 'manager'), ('veri', 'data_entry'), ('veri2', 'data_entry')]:
            u = User.objects.create_user(name, f'{name}@test.com', 'testpass123')
            UserProfile.objects.create(user=u, active_company=self.company)
            CompanyMembership.objects.create(company=self.company, user=u, role=role)
            self.users[name] = u
        self.client = APIClient()

    def _as(self, name):
        self.client.force_authenticate(user=self.users[name])

    def _target(self):
        return self.client.post('/api/emissions/targets/', {
            'title': '2030', 'base_year': 2024, 'target_year': 2030,
            'base_emissions_kg': 1000, 'target_reduction_percent': 30,
        }, format='json')

    def test_data_entry_cannot_change_targets_but_can_read_them(self):
        self._as('veri')
        self.assertEqual(self._target().status_code, 403)
        self._as('mudur')
        res = self._target()
        self.assertEqual(res.status_code, 201, res.content)
        tid = res.json()['id']
        self._as('veri')
        self.assertEqual(self.client.get('/api/emissions/targets/').status_code, 200)
        self.assertEqual(self.client.delete(f'/api/emissions/targets/{tid}/').status_code, 403)
        self.assertEqual(self.client.patch(f'/api/emissions/targets/{tid}/', {'title': 'x'}, format='json').status_code, 403)
        self._as('sahip')
        self.assertEqual(self.client.delete(f'/api/emissions/targets/{tid}/').status_code, 204)

    def _request(self):
        return self.client.post('/api/emissions/custom-requests/', {
            'scope': 'scope1', 'category_name': 'Jeneratör', 'source_name': 'Dizel jeneratör',
            'description': 'x', 'unit': 'litre', 'quantity': 120.5, 'year': 2026, 'month': 3,
        }, format='json')

    def test_custom_request_withdrawn_by_sender_not_by_other_member(self):
        self._as('veri')
        rid = self._request().json()['id']
        self._as('veri2')
        self.assertEqual(self.client.delete(f'/api/emissions/custom-requests/{rid}/').status_code, 403)
        self._as('veri')
        self.assertEqual(self.client.delete(f'/api/emissions/custom-requests/{rid}/').status_code, 204)

    def test_reviewed_custom_request_cannot_be_changed(self):
        from .models import CustomEmissionRequest
        self._as('veri')
        rid = self._request().json()['id']
        CustomEmissionRequest.objects.filter(id=rid).update(status='approved')
        self._as('sahip')
        res = self.client.delete(f'/api/emissions/custom-requests/{rid}/')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.json()['code'], 'already_reviewed')


class EntryChangeRulesTests(TestCase):
    """A data-entry member changes only their own entries; changing the amount
    of an approved entry sends it back for approval; every change is in the
    company's history."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        from accounts.models import UserProfile
        self.factor = EmissionFactor.objects.create(
            slug='test-gas-rules', name='Test Gas', name_tr='Test Gaz',
            scope='scope1', category='stationary_combustion', country='global',
            unit='m3', factor_kg_co2e=2, year=2024, source='generic',
            is_active=True, is_default=True,
        )
        self.company = Company.objects.create(
            legal_entity_name='Kural A.Ş.', tax_number='1234567892',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        self.users = {}
        for name, role in [('sahip', 'owner'), ('veri', 'data_entry'), ('denetci', 'auditor')]:
            u = User.objects.create_user(name, f'{name}@test.com', 'testpass123')
            UserProfile.objects.create(user=u, active_company=self.company)
            CompanyMembership.objects.create(company=self.company, user=u, role=role)
            self.users[name] = u
        self.client = APIClient()

    def _as(self, name):
        self.client.force_authenticate(user=self.users[name])

    def _entry(self, qty=300):
        res = self.client.post('/api/emissions/entries/', {
            'emission_factor': self.factor.id, 'year': 2026, 'month': 7, 'quantity': qty}, format='json')
        self.assertEqual(res.status_code, 201, res.content)
        return res.json()['id']

    def test_data_entry_cannot_change_or_delete_others_entries(self):
        self._as('sahip')
        eid = self._entry()
        self._as('veri')
        self.assertEqual(self.client.patch(f'/api/emissions/entries/{eid}/', {'quantity': 3}, format='json').status_code, 403)
        self.assertEqual(self.client.delete(f'/api/emissions/entries/{eid}/').status_code, 403)
        self.assertEqual(EmissionEntry.objects.get(id=eid).quantity, 300)

    def test_changing_an_approved_amount_goes_back_for_approval(self):
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved')
        res = self.client.patch(f'/api/emissions/entries/{eid}/', {'quantity': 3}, format='json')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertTrue(res.json()['pending_change'])
        # The approved value stays (and counts); the change waits next to it.
        old = EmissionEntry.objects.get(id=eid)
        self.assertEqual((old.status, old.quantity), ('approved', 300))
        change = EmissionEntry.objects.get(replaces_entry=old)
        self.assertEqual((change.status, change.quantity), ('submitted', 3))
        self.assertEqual(res.json()['replaces']['quantity'], 300.0)
        # A second edit replaces the waiting change, not adds another.
        self.client.patch(f'/api/emissions/entries/{eid}/', {'quantity': 4}, format='json')
        self.assertEqual(EmissionEntry.objects.filter(replaces_entry=old).count(), 1)
        # Approving the change puts it in the approved entry's place.
        change = EmissionEntry.objects.get(replaces_entry=old)
        self._as('sahip')
        self.client.post(f'/api/emissions/entries/{change.id}/approve/', {'action': 'approve'}, format='json')
        self.assertFalse(EmissionEntry.objects.filter(id=eid).exists())
        change.refresh_from_db()
        self.assertEqual((change.status, change.quantity, change.replaces_entry_id), ('approved', 4, None))
        # A description-only edit keeps the approval.
        self._as('veri')
        self.client.patch(f'/api/emissions/entries/{change.id}/', {'description': 'fatura no 12'}, format='json')
        self.assertEqual(EmissionEntry.objects.get(id=change.id).status, 'approved')

    def test_deleting_an_approved_entry_is_a_request(self):
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved')
        self.assertEqual(self.client.delete(f'/api/emissions/entries/{eid}/').status_code, 202)
        self.assertTrue(EmissionEntry.objects.get(id=eid).delete_requested)
        self._as('sahip')
        pending = self.client.get('/api/emissions/pending/').json()
        self.assertEqual([(p['id'], p['delete_requested']) for p in pending], [(eid, True)])
        # Refused: it stays.
        self.client.post(f'/api/emissions/entries/{eid}/approve/', {'action': 'reject'}, format='json')
        e = EmissionEntry.objects.get(id=eid)
        self.assertEqual((e.status, e.delete_requested), ('approved', False))
        # Asked again and accepted: deleted.
        self._as('veri')
        self.client.delete(f'/api/emissions/entries/{eid}/')
        self._as('sahip')
        self.assertEqual(self.client.post(f'/api/emissions/entries/{eid}/approve/', {'action': 'approve'},
                                          format='json').json()['status'], 'deleted')
        self.assertFalse(EmissionEntry.objects.filter(id=eid).exists())
        # A waiting entry of their own is still deleted at once.
        self._as('veri')
        wid = self._entry()
        self.assertEqual(self.client.delete(f'/api/emissions/entries/{wid}/').status_code, 204)

    def _approved_entry_with_change(self, qty=400):
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved')
        res = self.client.patch(f'/api/emissions/entries/{eid}/', {'quantity': qty}, format='json')
        self.assertTrue(res.json()['pending_change'])
        return eid, res.json()['id']

    def test_delete_request_drops_the_waiting_change(self):
        eid, cid = self._approved_entry_with_change()
        self.assertEqual(self.client.delete(f'/api/emissions/entries/{eid}/').status_code, 202)
        self.assertFalse(EmissionEntry.objects.filter(id=cid).exists())
        self._as('sahip')
        self.assertEqual([p['id'] for p in self.client.get('/api/emissions/pending/').json()], [eid])
        self.client.post(f'/api/emissions/entries/{eid}/approve/', {'action': 'approve'}, format='json')
        self.assertFalse(EmissionEntry.objects.filter(company=self.company).exists())
        # A change after a deletion request takes its place.
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved')
        self.client.delete(f'/api/emissions/entries/{eid}/')
        self.client.patch(f'/api/emissions/entries/{eid}/', {'quantity': 500}, format='json')
        self.assertFalse(EmissionEntry.objects.get(id=eid).delete_requested)

    def test_deleting_the_approved_entry_drops_its_waiting_change(self):
        eid, cid = self._approved_entry_with_change()
        self._as('sahip')
        self.assertEqual(self.client.delete(f'/api/emissions/entries/{eid}/').status_code, 204)
        self.assertFalse(EmissionEntry.objects.filter(id=cid).exists())

    def test_moving_an_approved_entry_to_another_facility_waits_for_approval(self):
        from companies.models import Facility
        a = Facility.objects.create(company=self.company, name='Fabrika A', country='TR')
        b = Facility.objects.create(company=self.company, name='Depo B', country='TR')
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved', facility=a)
        res = self.client.patch(f'/api/emissions/entries/{eid}/', {'facility': b.id}, format='json')
        self.assertTrue(res.json()['pending_change'])
        self.assertEqual(EmissionEntry.objects.get(id=eid).facility_id, a.id)
        self._as('sahip')
        row = self.client.get('/api/emissions/pending/').json()[0]
        self.assertEqual((row['replaces']['facility_name'], row['facility_name']), ('Fabrika A', 'Depo B'))

    def test_deletion_request_withdrawn_or_refused_with_reason(self):
        from accounts.models import Notification
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved')
        self.client.delete(f'/api/emissions/entries/{eid}/')
        self.assertEqual(self.client.post(f'/api/emissions/entries/{eid}/cancel-delete/').json()['status'], 'kept')
        self.assertFalse(EmissionEntry.objects.get(id=eid).delete_requested)
        self.client.delete(f'/api/emissions/entries/{eid}/')
        self._as('sahip')
        self.client.post(f'/api/emissions/entries/{eid}/approve/', {'action': 'reject', 'reason': 'Kayıt doğru'},
                         format='json')
        n = Notification.objects.filter(user=self.users['veri']).latest('id')
        self.assertIn('Kayıt doğru', n.message)
        self.assertEqual(EmissionEntry.objects.get(id=eid).status, 'approved')

    def test_rejected_change_says_the_approved_figure_stays(self):
        from accounts.models import Notification
        eid, cid = self._approved_entry_with_change()
        self._as('sahip')
        self.client.post(f'/api/emissions/entries/{cid}/approve/', {'action': 'reject', 'reason': 'Fatura'},
                         format='json')
        n = Notification.objects.filter(user=self.users['veri']).latest('id')
        self.assertIn('300', n.message)

    def test_approvers_are_told_a_change_is_a_change(self):
        from accounts.models import Notification
        self._approved_entry_with_change(qty=250)
        n = Notification.objects.filter(user=self.users['sahip']).latest('id')
        self.assertEqual(n.title, 'Değişiklik talebi')
        self.assertIn('300 m³ → 250 m³', n.message)
        # The link names the entry, so the review page opens on it.
        self.assertRegex(n.link, r'^/dashboard\?tab=review&entry=\d+$')

    def test_export_marks_changes_and_deletion_requests(self):
        from .views import _export_rows
        eid, cid = self._approved_entry_with_change(qty=250)
        other = self._entry(qty=50)
        EmissionEntry.objects.filter(id=other).update(status='approved', delete_requested=True)
        _, rows = _export_rows(EmissionEntry.objects.filter(company=self.company).order_by('id'), 'tr')
        self.assertEqual([r[12] for r in rows],
                         ['Onaylı', 'Değişiklik onayda (onaylı: 300 m³)', 'Onaylı · silme talebi'])

    def test_only_the_author_withdraws_a_deletion_request(self):
        self._as('veri')
        eid = self._entry()
        EmissionEntry.objects.filter(id=eid).update(status='approved')
        self.client.delete(f'/api/emissions/entries/{eid}/')
        self._as('sahip')
        res = self.client.post(f'/api/emissions/entries/{eid}/cancel-delete/')
        self.assertEqual((res.status_code, res.json()['code']), (403, 'not_requester'))
        self.assertTrue(EmissionEntry.objects.get(id=eid).delete_requested)

    def test_waiting_changes_are_counted_apart(self):
        from .inventory import not_counted, not_counted_note
        eid, cid = self._approved_entry_with_change()
        self._entry(qty=50)
        info = not_counted(EmissionEntry.objects.filter(company=self.company, year=2026))
        self.assertEqual(info['pending'], {'count': 1, 'total_kg': 100.0})
        self.assertEqual(info['changes'], {'count': 1})
        self.assertIn('1 değişiklik talebi', not_counted_note(info, tr=True))

    def test_owner_change_stays_approved_and_is_in_history(self):
        self._as('sahip')
        eid = self._entry()
        self.assertEqual(EmissionEntry.objects.get(id=eid).status, 'approved')
        self.client.patch(f'/api/emissions/entries/{eid}/', {'quantity': 250}, format='json')
        self.assertEqual(EmissionEntry.objects.get(id=eid).status, 'approved')
        self._as('denetci')
        res = self.client.get('/api/accounts/history/')
        self.assertEqual(res.status_code, 200)
        actions = [r['action'] for r in res.json()]
        self.assertEqual(actions[:2], ['entry_updated', 'entry_created'])
        # Shown in the reader's language (accounts/history_text.py).
        self.assertIn('300 m³ → Test Gaz · Temmuz 2026 · 250 m³', res.json()[0]['detail'])

    def test_data_entry_cannot_read_history(self):
        self._as('veri')
        self.assertEqual(self.client.get('/api/accounts/history/').status_code, 403)


class ProofDocumentTests(EntryChangeRulesTests):
    """Proof can be attached, replaced and removed after an entry exists, and
    a proof file lost from storage answers 410 instead of a 500."""

    def _pdf(self, name='fatura.pdf'):
        from django.core.files.uploadedfile import SimpleUploadedFile
        return SimpleUploadedFile(name, b'%PDF-1.4 test', content_type='application/pdf')

    def setUp(self):
        super().setUp()
        import tempfile
        from django.test import override_settings
        self._media = tempfile.TemporaryDirectory()
        self._override = override_settings(MEDIA_ROOT=self._media.name)
        self._override.enable()
        self.addCleanup(self._override.disable)
        self.addCleanup(self._media.cleanup)

    def test_attach_replace_and_remove_proof_later(self):
        self._as('veri')
        eid = self._entry()
        res = self.client.post(f'/api/emissions/entries/{eid}/proof/', {'proof_document': self._pdf()}, format='multipart')
        self.assertEqual(res.status_code, 200, res.content)
        self.assertTrue(res.json()['proof_available'])
        self.assertEqual(self.client.get(f'/api/emissions/entries/{eid}/proof/').status_code, 200)
        res = self.client.post(f'/api/emissions/entries/{eid}/proof/', {'proof_document': self._pdf('yeni.pdf')}, format='multipart')
        self.assertIn('yeni', res.json()['proof_document'])
        self.assertEqual(self.client.delete(f'/api/emissions/entries/{eid}/proof/').status_code, 204)
        self.assertFalse(EmissionEntry.objects.get(id=eid).proof_document)

    def test_other_members_entry_proof_is_protected(self):
        self._as('sahip')
        eid = self._entry()
        self._as('veri')
        res = self.client.post(f'/api/emissions/entries/{eid}/proof/', {'proof_document': self._pdf()}, format='multipart')
        self.assertEqual(res.status_code, 403)

    def test_bad_file_type_is_refused(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        self._as('veri')
        eid = self._entry()
        res = self.client.post(f'/api/emissions/entries/{eid}/proof/',
                               {'proof_document': SimpleUploadedFile('x.exe', b'MZ')}, format='multipart')
        self.assertEqual(res.status_code, 400)

    def test_lost_file_answers_410(self):
        self._as('veri')
        eid = self._entry()
        self.client.post(f'/api/emissions/entries/{eid}/proof/', {'proof_document': self._pdf()}, format='multipart')
        entry = EmissionEntry.objects.get(id=eid)
        entry.proof_document.storage.delete(entry.proof_document.name)
        res = self.client.get(f'/api/emissions/entries/{eid}/proof/')
        self.assertEqual(res.status_code, 410)
        self.assertEqual(res.json()['code'], 'proof_missing')
        listed = self.client.get('/api/emissions/entries/?year=2026').json()
        listed = listed.get('results', listed)
        self.assertFalse([e for e in listed if e['id'] == eid][0]['proof_available'])


class OnlyApprovedEntriesCountTests(TestCase):
    """Totals and reports count approved entries only; pending and rejected
    ones are reported separately, for information."""

    def setUp(self):
        from companies.models import Company, CompanyMembership
        self.user = User.objects.create_user('countu', 'count@test.com', 'testpass123')
        self.company = Company.objects.create(
            legal_entity_name='Sayım A.Ş.', tax_number='7777777777',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1,
        )
        CompanyMembership.objects.create(company=self.company, user=self.user, role='owner')
        self.factor = EmissionFactor.objects.create(
            slug='count-gas', name='Count Gas', scope='scope1', category='stationary_combustion',
            country='global', unit='kg', factor_kg_co2e=2, year=2024, source='generic',
            is_active=True, is_default=True,
        )
        for qty, status in ((100, 'approved'), (40, 'submitted'), (10, 'draft')):
            EmissionEntry.objects.create(user=self.user, company=self.company, emission_factor=self.factor,
                                         year=2026, month=3, quantity=qty, status=status)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def test_summary_counts_approved_only_and_reports_the_rest(self):
        data = self.client.get('/api/emissions/summary/?year=2026').data
        self.assertAlmostEqual(data['total_kg'], 200.0)
        self.assertAlmostEqual(data['scope1_kg'], 200.0)
        self.assertEqual(data['not_counted']['pending'], {'count': 1, 'total_kg': 80.0})
        self.assertEqual(data['not_counted']['rejected'], {'count': 1, 'total_kg': 20.0})

    def test_year_comparison_and_facility_breakdown_count_approved_only(self):
        data = self.client.get('/api/emissions/comparison/?year1=2025&year2=2026').data
        self.assertAlmostEqual(data['year2']['total_kg'], 200.0)
        from companies.models import Facility
        fac = Facility.objects.create(company=self.company, name='Fab', country='TR')
        EmissionEntry.objects.filter(company=self.company).update(facility=fac)
        rows = self.client.get('/api/emissions/by-facility/?year=2026').data
        self.assertAlmostEqual(rows[0]['total_kg'], 200.0)

    def test_note_text(self):
        from .inventory import not_counted, not_counted_note
        info = not_counted(EmissionEntry.objects.filter(company=self.company, year=2026))
        tr = not_counted_note(info, tr=True)
        self.assertIn('1 onay bekleyen kayıt (0,080 tCO₂e)', tr)
        self.assertIn('1 reddedilen kayıt (0,020 tCO₂e)', tr)
        self.assertIn('dahil değildir', tr)
        self.assertEqual(not_counted_note({'pending': {'count': 0}, 'rejected': {'count': 0}}), '')

    def test_emissions_pdf_carries_the_note(self):
        from io import BytesIO
        from pypdf import PdfReader
        res = self.client.get('/api/emissions/report/?year=2026&lang=tr')
        self.assertEqual(res.status_code, 200)
        text = ' '.join(' '.join(p.extract_text() for p in PdfReader(BytesIO(res.content)).pages).split())
        self.assertIn('Toplam emisyon: 0,20 tCO₂e', text)  # approved only
        self.assertIn('onay bekleyen kayıt', text)
        self.assertIn('dahil değildir', text)

    def test_legacy_migration_approves_only_approver_entries(self):
        import importlib
        from django.apps import apps
        from companies.models import CompanyMembership
        member = User.objects.create_user('dataent', 'de@test.com', 'testpass123')
        CompanyMembership.objects.create(company=self.company, user=member, role='data_entry')
        mine = EmissionEntry.objects.filter(company=self.company, status='submitted').get()
        theirs = EmissionEntry.objects.create(user=member, company=self.company, emission_factor=self.factor,
                                              year=2026, month=4, quantity=5, status='submitted')
        rejected = EmissionEntry.objects.filter(company=self.company, status='draft').get()
        mig = importlib.import_module('emissions.migrations.0016_approve_legacy_approver_entries')
        mig.approve_legacy(apps, None)
        mine.refresh_from_db(); theirs.refresh_from_db(); rejected.refresh_from_db()
        self.assertEqual(mine.status, 'approved')
        self.assertEqual(theirs.status, 'submitted')
        self.assertEqual(rejected.status, 'draft')


class CompanyCountryFactorTests(TestCase):
    """The AI chat / questionnaire factor lookup follows the company's HQ:
    Turkish factors for a Turkish company, global ones and the country's own
    grid for a company based elsewhere."""

    def setUp(self):
        from companies.models import Company
        mk = lambda slug, country, value, **kw: EmissionFactor.objects.create(
            slug=slug, name=slug, scope=kw.get('scope', 'scope2'), category=kw.get('category', 'electricity'),
            country=country, unit=kw.get('unit', 'kwh'), factor_kg_co2e=value, year=2024,
            source='generic', is_active=True, is_default=True)
        mk('turkey-grid', 'turkey', 0.42)
        mk('uk-grid', 'global', 0.207)
        mk('eu-grid', 'global', 0.295)
        mk('grid-average', 'global', 0.475)
        mk('natural-gas-m3', 'turkey', 2.02, scope='scope1', category='stationary_combustion', unit='m3')
        mk('natural-gas-m3', 'global', 2.03, scope='scope1', category='stationary_combustion', unit='m3')
        base = dict(tax_number='', countries_of_operation='x', main_activity_description='x',
                    number_of_employees='1-10', annual_turnover_range='x', number_of_facilities=1)
        self.tr = Company.objects.create(legal_entity_name='TR', country_of_headquarters='TR', **base)
        self.gb = Company.objects.create(legal_entity_name='GB', country_of_headquarters='GB', **base)
        self.nl = Company.objects.create(legal_entity_name='NL', country_of_headquarters='NL', **base)
        self.br = Company.objects.create(legal_entity_name='BR', country_of_headquarters='BR', **base)

    def _slug(self, activity, unit, company):
        from .factor_lookup import resolve_factor
        factor, _u, _m, err = resolve_factor(activity, unit, company)
        self.assertIsNone(err)
        return factor.slug, factor.country

    def test_electricity_grid_follows_the_company_country(self):
        self.assertEqual(self._slug('electricity', 'kwh', self.tr), ('turkey-grid', 'turkey'))
        self.assertEqual(self._slug('electricity', 'kwh', None), ('turkey-grid', 'turkey'))
        self.assertEqual(self._slug('electricity', 'kwh', self.gb), ('uk-grid', 'global'))
        self.assertEqual(self._slug('electricity', 'kwh', self.nl), ('eu-grid', 'global'))
        self.assertEqual(self._slug('electricity', 'kwh', self.br), ('grid-average', 'global'))  # no brazil-grid here

    def test_grid_matches_country(self):
        from .factor_lookup import grid_matches_country
        f = lambda slug: EmissionFactor.objects.filter(slug=slug).first()
        self.assertTrue(grid_matches_country(f('turkey-grid'), 'tr'))
        self.assertFalse(grid_matches_country(f('turkey-grid'), 'DE'))
        self.assertTrue(grid_matches_country(f('uk-grid'), 'GB'))
        self.assertFalse(grid_matches_country(f('uk-grid'), 'AT'))
        self.assertTrue(grid_matches_country(f('eu-grid'), 'AT'))
        self.assertFalse(grid_matches_country(f('eu-grid'), 'TR'))
        self.assertIsNone(grid_matches_country(f('grid-average'), 'TR'))
        self.assertIsNone(grid_matches_country(f('uk-grid'), ''))
        self.assertIsNone(grid_matches_country(f('natural-gas-m3'), 'TR'))

    def test_entry_shows_a_factor_of_another_country(self):
        # A Turkish company's site in Germany on the Turkish grid is marked;
        # a site in Turkey is not. The calculation is unchanged.
        from django.contrib.auth.models import User
        from companies.models import Facility
        from .models import EmissionEntry
        from .serializers import EmissionEntrySerializer
        user = User.objects.create_user('fc', 'fc@t.com', 'x')
        grid = EmissionFactor.objects.get(slug='turkey-grid')
        de = Facility.objects.create(company=self.tr, name='Werk Stuttgart', country='DE')
        tr = Facility.objects.create(company=self.tr, name='Gebze', country='TR')
        mk = lambda fac: EmissionEntry.objects.create(user=user, company=self.tr, emission_factor=grid,
                                                      year=2025, month=1, quantity=100, facility=fac)
        data = EmissionEntrySerializer(mk(de)).data
        self.assertTrue(data['factor_country_mismatch'])
        self.assertEqual(data['facility_country'], 'DE')
        self.assertFalse(EmissionEntrySerializer(mk(tr)).data['factor_country_mismatch'])
        self.assertFalse(EmissionEntrySerializer(mk(None)).data['factor_country_mismatch'])

    def test_iea_source_has_a_name(self):
        # Shown as "Registered factor — iea_2024" in the chat before.
        self.assertEqual(EmissionFactor(source='iea_2024').get_source_display(), 'IEA 2024')

    def test_foreign_company_prefers_the_global_version(self):
        self.assertEqual(self._slug('natural_gas', 'm3', self.tr), ('natural-gas-m3', 'turkey'))
        self.assertEqual(self._slug('natural_gas', 'm3', self.gb), ('natural-gas-m3', 'global'))


class CommercialFlightScopeTests(TestCase):
    """Commercial flights are Scope 3 Cat. 6 business travel (GHG Protocol),
    in the seed data and after the migration."""

    def test_seed_data(self):
        from . import seed_data
        rows = [r for rows in (v for v in vars(seed_data).values() if isinstance(v, list)) for r in rows
                if isinstance(r, dict) and r.get('slug') in ('flight-domestic', 'flight-international')]
        self.assertEqual(len(rows), 2)
        for r in rows:
            self.assertEqual((r['scope'], r['category']), ('scope3', 'business_travel'))

    def test_migration_moves_existing_rows(self):
        import importlib
        from django.apps import apps
        f = EmissionFactor.objects.create(slug='flight-domestic', name='Domestic Flight (Turkey)', scope='scope1',
                                          category='mobile_combustion', country='turkey', unit='km',
                                          factor_kg_co2e=0.232, year=2025, source='turkey_fleet',
                                          is_active=True, is_default=True)
        importlib.import_module('emissions.migrations.0017_commercial_flights_scope3').move_to_scope3(apps, None)
        f.refresh_from_db()
        self.assertEqual((f.scope, f.category, float(f.factor_kg_co2e)), ('scope3', 'business_travel', 0.232))


class PublicTransportScopeTests(TestCase):
    """Turkish public transport (train, metro, bus, dolmuş) is Scope 3 Cat. 6
    business travel, like commercial flights."""
    SLUGS = ('train', 'metro', 'bus', 'dolmus')

    def test_seed_data(self):
        from . import seed_data
        rows = [r for rows in (v for v in vars(seed_data).values() if isinstance(v, list)) for r in rows
                if isinstance(r, dict) and r.get('slug') in self.SLUGS and r.get('country') == 'turkey']
        self.assertEqual(len(rows), 4)
        for r in rows:
            self.assertEqual((r['scope'], r['category']), ('scope3', 'business_travel'))

    def test_migration_moves_existing_rows(self):
        import importlib
        from django.apps import apps
        for slug in self.SLUGS:
            EmissionFactor.objects.create(slug=slug, name=slug, scope='scope1', category='mobile_combustion',
                                          country='turkey', unit='km', factor_kg_co2e=0.05, year=2025,
                                          source='turkey_fleet', is_active=True, is_default=True)
        importlib.import_module('emissions.migrations.0018_public_transport_scope3').move_to_scope3(apps, None)
        for f in EmissionFactor.objects.filter(slug__in=self.SLUGS):
            self.assertEqual((f.scope, f.category, float(f.factor_kg_co2e)), ('scope3', 'business_travel', 0.05))


class FuturePeriodTests(ApprovalNotificationAndDuplicateTests):
    """No entry, chat save or custom request for a month that hasn't started."""

    def _next_month(self):
        from datetime import date
        t = date.today()
        return (t.year + 1, 1) if t.month == 12 else (t.year, t.month + 1)

    def test_helper(self):
        from datetime import date
        from .periods import is_future_period
        today = date(2026, 9, 29)
        self.assertFalse(is_future_period(2026, 9, today))
        self.assertFalse(is_future_period(2025, 12, today))
        self.assertTrue(is_future_period(2026, 10, today))
        self.assertTrue(is_future_period(2027, 1, today))

    def test_form_rejects_next_month_but_accepts_this_month(self):
        from datetime import date
        self._as('aylin')
        y, m = self._next_month()
        res = self._entry(year=y, month=m)
        self.assertEqual(res.status_code, 400)
        self.assertIn('future_period', str(res.data['code']))
        t = date.today()
        self.assertEqual(self._entry(year=t.year, month=t.month).status_code, 201)

    def test_moving_an_entry_into_the_future_is_rejected(self):
        self._as('aylin')
        entry_id = self._entry().data['id']
        y, m = self._next_month()
        res = self.client.patch(f'/api/emissions/entries/{entry_id}/', {'year': y, 'month': m}, format='json')
        self.assertEqual(res.status_code, 400)
        # Changing only the amount still works.
        self.assertEqual(self.client.patch(f'/api/emissions/entries/{entry_id}/', {'quantity': 11},
                                           format='json').status_code, 200)

    def test_chat_confirm_rejects_future_month_in_the_users_language(self):
        self._as('aylin')
        y, m = self._next_month()
        res = self.client.post('/api/chat/confirm-entry/', {
            'fuel_type': 'test-gas-dup', 'quantity': 10, 'unit': 'kg',
            'year': y, 'month': m, 'language': 'tr',
        }, format='json')
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data['code'], 'future_period')
        self.assertIn('henüz başlamadı', res.data['error'])
        self.assertEqual(EmissionEntry.objects.count(), 0)

    def test_custom_request_rejects_future_month(self):
        self._as('aylin')
        y, m = self._next_month()
        res = self.client.post('/api/emissions/custom-requests/', {
            'scope': 'scope1', 'category_name': 'Jeneratör', 'source_name': 'Dizel',
            'description': 'x', 'unit': 'litre', 'quantity': 5, 'year': y, 'month': m,
        }, format='json')
        self.assertEqual(res.status_code, 400)


class CustomRequestUnitTests(ApprovalNotificationAndDuplicateTests):
    def test_unit_must_be_short(self):
        self._as('aylin')
        data = {'scope': 'scope1', 'category_name': 'Jeneratör', 'source_name': 'Dizel', 'description': 'x',
                'quantity': 5, 'year': 2026, 'month': 1}
        res = self.client.post('/api/emissions/custom-requests/', {**data, 'unit': 'Soğutma gazı R-410A dolumu'},
                               format='json')
        self.assertEqual(res.status_code, 400)
        self.assertIn('unit', res.data)
        self.assertEqual(self.client.post('/api/emissions/custom-requests/', {**data, 'unit': 'litre'},
                                          format='json').status_code, 201)


class EntryYearBoundTests(TestCase):
    def test_year_before_1990_is_refused(self):
        from rest_framework.test import APIClient
        from companies.models import Company, CompanyMembership
        from emissions.models import EmissionFactor
        user = User.objects.create_user('yr', 'yr@test.com', 'pass12345')
        company = Company.objects.create(
            legal_entity_name='Yr Co', tax_number='12', country_of_headquarters='TR',
            countries_of_operation='TR', nace_code='', main_activity_description='x',
            number_of_employees='1-10', annual_turnover_range='x', number_of_facilities=1)
        CompanyMembership.objects.create(user=user, company=company, role='owner')
        f = EmissionFactor.objects.create(slug='yr-grid', country='turkey', year=2024, name='grid',
                                          scope='scope2', category='electricity', unit='kwh',
                                          factor_kg_co2e=0.4, source='test')
        c = APIClient(); c.force_authenticate(user=user)
        r = c.post('/api/emissions/entries/', {'emission_factor': f.id, 'quantity': 10, 'year': 1000, 'month': 1}, format='json')
        self.assertEqual(r.status_code, 400)
        self.assertIn('invalid_year', str(r.data))
        r = c.post('/api/emissions/entries/', {'emission_factor': f.id, 'quantity': 10, 'year': 2024, 'month': 1}, format='json')
        self.assertEqual(r.status_code, 201)


class ApprovalEmailTests(ApprovalNotificationAndDuplicateTests):
    """Approvers are e-mailed about waiting entries at most once a day; an
    author is e-mailed when an entry is rejected."""

    def test_one_email_a_day_for_waiting_entries(self):
        from django.core import mail
        self._as('ali')
        self._entry()
        self._entry(month=4)
        to_owner = [m for m in mail.outbox if m.to == ['aylin@test.com']]
        self.assertEqual(len(to_owner), 1)
        self.assertIn('onayınızı bekleyen kayıtlar var', to_owner[0].subject)
        self.assertIn('en az 1 kayıt', to_owner[0].body)
        self.assertIn('Kaya Tekstil', to_owner[0].body)

    def test_email_respects_the_preference(self):
        from django.core import mail
        from accounts.models import UserProfile
        UserProfile.objects.filter(user=self.users['aylin']).update(notify_approvals=False)
        self._as('ali')
        self._entry()
        self.assertEqual([m for m in mail.outbox if m.to == ['aylin@test.com']], [])

    def test_rejection_is_emailed_with_the_reason(self):
        from django.core import mail
        self._as('ali')
        entry_id = self._entry().data['id']
        mail.outbox.clear()
        self._as('aylin')
        self.client.post(f'/api/emissions/entries/{entry_id}/approve/',
                         {'action': 'reject', 'reason': 'Invoice missing'}, format='json')
        self.assertEqual(len(mail.outbox), 1)
        self.assertEqual(mail.outbox[0].to, ['ali@test.com'])
        self.assertIn('Invoice missing', mail.outbox[0].body)


class YearlySourceAndYearChangeTests(ApprovalNotificationAndDuplicateTests):
    """A yearly source entered for a second month is flagged; the summary
    compares a year with the one before."""

    def test_yearly_source_second_month_warns(self):
        f = EmissionFactor.objects.create(
            slug='test-franchise', name='Franchise operations (per franchise/year)',
            name_tr='Franchise operasyonları (yıllık)', scope='scope3', category='franchises',
            country='global', unit='franchises', factor_kg_co2e=15000, year=2024, source='generic',
            is_active=True)
        self._as('aylin')
        mk = lambda month, **x: self.client.post('/api/emissions/entries/', {
            'emission_factor': f.id, 'year': 2025, 'month': month, 'quantity': 2, **x}, format='json')
        self.assertEqual(mk(1).status_code, 201)
        r = mk(2)
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.data['code'], 'yearly_repeat')
        self.assertIn('Ocak', r.data['message_tr'])
        self.assertEqual(mk(2, confirm_duplicate=True).status_code, 201)
        # an ordinary monthly source is not flagged
        self._entry(month=1, year=2025)
        self.assertEqual(self._entry(month=2, year=2025, quantity=11).status_code, 201)

    def test_summary_compares_with_the_year_before(self):
        from datetime import date
        from .inventory import previous_year_change
        self._as('aylin')
        self._entry(year=2024, month=3, quantity=100)
        self._entry(year=2025, month=3, quantity=80)
        c = previous_year_change(self.company, 2025, 80 * 2.5, today=date(2026, 10, 11))
        self.assertEqual((c['previous_year'], c['change_pct'], c['months']), (2024, -20.0, 12))
        # running year: the same months only
        self._entry(year=2025, month=11, quantity=40)
        self._entry(year=2026, month=3, quantity=50)
        c = previous_year_change(self.company, 2026, 125, today=date(2026, 10, 11))
        self.assertEqual((c['months'], c['previous_kg']), (10, 200.0))
        self.assertIsNone(previous_year_change(self.company, 2024, 250, today=date(2026, 10, 11)))
        self.assertIn('previous_year_change', self.client.get('/api/emissions/summary/?year=2025').data)


class ListWorkspaceEntriesCommandTests(ApprovalNotificationAndDuplicateTests):
    """The read-only listing of old Data Panel entries flags the doubles."""

    def test_lists_data_panel_entries_and_flags_doubles(self):
        from io import StringIO
        from django.core.management import call_command
        mk = lambda desc, year: EmissionEntry.objects.create(
            user=self.users['aylin'], company=self.company, emission_factor=self.factor,
            year=year, month=1, quantity=1000, calculated_co2e_kg=2500,
            description=desc, status='approved')
        mk('Questionnaire step 3A-5', 2023)
        doubled = mk('Workspace 3A (report 14)', 2023)
        mk('Workspace 3A (report 9)', 2022)
        out = StringIO()
        call_command('list_workspace_entries', stdout=out)
        text = out.getvalue()
        self.assertIn(f'#{doubled.id}', text)
        self.assertIn('2 Data Panel entries, 1 likely counted twice.', text)
        self.assertEqual(EmissionEntry.objects.count(), 3)


class AmountCeilingTests(TestCase):
    """Fix 1: an over-capacity amount is rejected cleanly (no 500), while a
    realistically-large amount is still accepted and calculates co2e."""

    def setUp(self):
        self.user = User.objects.create_user('ceiluser', 'ceil@test.com', 'testpass123')
        from companies.models import Company, CompanyMembership
        self.company = Company.objects.create(
            legal_entity_name='Ceiling Co', tax_number='999',
            country_of_headquarters='TR', countries_of_operation='TR',
            main_activity_description='Test', number_of_employees='1-10',
            annual_turnover_range='<1M', number_of_facilities=1,
        )
        CompanyMembership.objects.create(company=self.company, user=self.user, role='owner')
        # Resolver looks up electricity/kwh via slug 'turkey-grid'.
        self.factor = EmissionFactor.objects.create(
            slug='turkey-grid', name='Turkey Grid', name_tr='Türkiye Şebeke',
            scope='scope2', category='electricity', country='turkey',
            unit='kwh', factor_kg_co2e=Decimal('0.4199'), year=2023, source='atom_kablo',
            is_active=True, is_default=True,
        )

    def test_resolver_rejects_over_capacity(self):
        from .factor_lookup import resolve_factor_and_amount, AMOUNT_TOO_LARGE_ERROR
        factor, qty, co2e, err = resolve_factor_and_amount('electricity', 999000000000, 'kwh', None)
        self.assertIsNone(factor)
        self.assertIsNone(qty)
        self.assertIsNone(co2e)
        self.assertEqual(err, AMOUNT_TOO_LARGE_ERROR)

    def test_resolver_accepts_realistic_value(self):
        from .factor_lookup import resolve_factor_and_amount
        factor, qty, co2e, err = resolve_factor_and_amount('electricity', 5000000, 'kwh', None)
        self.assertIsNone(err)
        self.assertIsNotNone(factor)
        self.assertGreater(co2e, 0)

    def test_create_entry_rejects_over_capacity_without_row(self):
        from .factor_lookup import create_entry_from_activity, AMOUNT_TOO_LARGE_ERROR
        before = EmissionEntry.objects.count()
        entry, err = create_entry_from_activity(
            self.user, self.company, 'electricity', 999000000000, 'kwh', 2024, 1)
        self.assertIsNone(entry)
        self.assertEqual(err, AMOUNT_TOO_LARGE_ERROR)
        self.assertEqual(EmissionEntry.objects.count(), before)

    def test_create_entry_accepts_realistic_value(self):
        from .factor_lookup import create_entry_from_activity
        entry, err = create_entry_from_activity(
            self.user, self.company, 'electricity', 5000000, 'kwh', 2024, 1)
        self.assertIsNone(err)
        self.assertIsNotNone(entry)
        self.assertGreater(float(entry.calculated_co2e_kg), 0)


class ApproverSeesOverlapTests(ApprovalNotificationAndDuplicateTests):
    """A monthly entry of a source the year's inventory already holds reaches
    the approver with the same warning its author was shown."""

    def test_pending_entry_carries_the_overlap_warning(self):
        EmissionEntry.objects.create(
            user=self.users['aylin'], company=self.company, emission_factor=self.factor,
            year=2025, month=1, quantity=1000, calculated_co2e_kg=2500,
            description='Questionnaire step 3A-5', status='approved')
        self._as('ali')
        r = self._entry(year=2025, month=9)
        self.assertEqual(r.status_code, 409)
        self.assertEqual(r.data['code'], 'annual_overlap')
        r = self._entry(year=2025, month=9, confirm_duplicate=True)
        self.assertEqual(r.status_code, 201, r.content)
        self._as('aylin')
        rows = self.client.get('/api/emissions/pending/').data
        row = next(x for x in rows if x['id'] == r.data['id'])
        self.assertIn('iki kez sayılabilir', row['overlap_warning']['tr'])
        self.assertIn('counted twice', row['overlap_warning']['en'])

    def test_rejection_reason_full_stop_is_not_doubled(self):
        from accounts.models import Notification
        self._as('ali')
        entry_id = self._entry(year=2025, month=9).data['id']
        self._as('aylin')
        r = self.client.post(f'/api/emissions/entries/{entry_id}/approve/', {'action': 'reject', 'reason': 'Fatura eksik.'}, format='json')
        self.assertIn(r.status_code, (200, 204), r.content)
        n = Notification.objects.filter(user=self.users['ali'], notification_type='entry_rejected').latest('id')
        self.assertNotIn('..', n.message)
        self.assertTrue(n.link.endswith(f'&entry={entry_id}'))


class CustomRequestReviewTests(ApprovalNotificationAndDuplicateTests):
    """A reviewed custom request reaches its author in their language and
    number format; the chat names approved requests instead of "no data"."""

    def _request(self, **extra):
        from emissions.models import CustomEmissionRequest
        return CustomEmissionRequest.objects.create(
            user=self.users['ali'], company=self.company, scope='scope1',
            category_name='Jeneratör', source_name='Biyodizel jeneratör', description='x',
            unit='litre', quantity=1250.5, year=2026, month=9, **extra)

    def _review(self, req, status, **fields):
        from django.contrib.admin.sites import site
        from emissions.admin import CustomEmissionRequestAdmin
        from emissions.models import CustomEmissionRequest
        req.status = status
        for k, v in fields.items():
            setattr(req, k, v)

        class Form:
            changed_data = ['status']
        CustomEmissionRequestAdmin(CustomEmissionRequest, site).save_model(None, req, Form(), True)

    def test_approval_is_in_the_authors_language_and_opens_the_year(self):
        req = self._request()
        self._review(req, 'approved', approved_factor_kg_co2e=0.5)
        [note] = self._notes('ali', 'custom_approved')  # ali reads English
        self.assertEqual(note.title, 'Your custom request was approved')
        self.assertIn('1,250.5 litre · September 2026', note.message)
        self.assertIn('625.25 kg CO₂e (0.63 t)', note.message)
        self.assertEqual(note.link, '/dashboard?tab=emissions&year=2026')

    def test_turkish_author_gets_turkish_numbers_and_the_reason(self):
        profile = self.users['ali'].profile
        profile.language_preference = 'tr'
        profile.save()
        req = self._request()
        self._review(req, 'rejected', admin_notes='Fatura eksik.')
        [note] = self._notes('ali', 'custom_rejected')
        self.assertEqual(note.title, 'Özel talebiniz reddedildi')
        self.assertIn('1.250,5 litre · Eylül 2026 reddedildi. Neden: Fatura eksik.', note.message)

    def test_chat_total_names_approved_requests(self):
        from chat.local_answers import _total_answer
        self.assertIn('kayıtlı emisyon verisi yok', _total_answer(self.company, 2026, 'tr'))
        self._review(self._request(), 'approved', approved_factor_kg_co2e=0.5)
        answer = _total_answer(self.company, 2026, 'tr')
        self.assertIn('onaylı emisyon kaydı yok', answer)
        self.assertIn('Onaylı özel talepleriniz: **0,63 tCO₂e**', answer)

    def test_request_keeps_a_facility_of_the_company_only(self):
        from companies.models import Company, Facility
        own = Facility.objects.create(company=self.company, name='Gebze')
        other_co = Company.objects.create(
            legal_entity_name='Other', tax_number='9', country_of_headquarters='TR',
            countries_of_operation='TR', main_activity_description='x', number_of_employees='1-10',
            annual_turnover_range='x', number_of_facilities=1)
        other = Facility.objects.create(company=other_co, name='Elsewhere')
        self._as('ali')
        data = {'scope': 'scope1', 'category_name': 'Jeneratör', 'source_name': 'Dizel',
                'description': 'x', 'unit': 'litre', 'quantity': 5, 'year': 2026, 'month': 1}
        res = self.client.post('/api/emissions/custom-requests/', {**data, 'facility': own.id}, format='json')
        self.assertEqual(res.status_code, 201)
        self.assertEqual(res.data['facility'], own.id)
        res = self.client.post('/api/emissions/custom-requests/', {**data, 'facility': other.id}, format='json')
        self.assertEqual(res.status_code, 400)
