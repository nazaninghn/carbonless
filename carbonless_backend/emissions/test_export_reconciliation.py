"""F-2 (Stage 1 audit): the exports reconcile with the Dashboard.

The Dashboard total is the approved entries plus the approved custom requests that are
counted on their own (status approved, an amount calculated, not a facility electricity
factor). The CSV and Excel exports listed entries only, so a company's export summed to less
than its Dashboard. These tests use one isolated company and compare the Dashboard's figures
with the rows of each export, row by row (inclusion, quantity, unit, factor, kg, tonnes)."""
import csv
import io
from decimal import Decimal

from django.contrib.auth.models import User
from django.test import TestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from accounts.models import UserProfile
from companies.models import Company, CompanyMembership, Facility
from .models import CustomEmissionRequest, EmissionEntry, EmissionFactor

YEAR = 2026


def _company(name, tax):
    return Company.objects.create(
        legal_entity_name=name, tax_number=tax, country_of_headquarters='TR', countries_of_operation='TR',
        main_activity_description='x', number_of_employees='1-10', annual_turnover_range='x', number_of_facilities=1)


class ExportReconcilesWithDashboardTests(TestCase):

    def setUp(self):
        self.company = _company('Recon Co', '700')
        self.user = User.objects.create_user('recon', 'recon@test.com', 'testpass123')
        UserProfile.objects.create(user=self.user, active_company=self.company)
        CompanyMembership.objects.create(company=self.company, user=self.user, role='owner')
        self.facility = Facility.objects.create(company=self.company, name='Gebze', country='TR', city='Kocaeli')
        factor = EmissionFactor.objects.create(
            slug='recon-gas', name='Natural Gas', name_tr='Doğal Gaz', scope='scope1',
            category='stationary_combustion', country='global', unit='gj', factor_kg_co2e=56.211,
            year=2024, source='generic', is_active=True, is_default=True)
        mk = lambda status, qty, month: EmissionEntry.objects.create(
            user=self.user, company=self.company, emission_factor=factor, year=YEAR, month=month,
            quantity=qty, status=status)
        mk('approved', 10, 3)
        mk('approved', 4, 5)
        mk('submitted', 99, 6)                      # pending: listed, not in any total
        self.other = _company('Other Co', '701')
        self.counted = [
            self._request(scope='scope1', source='Forklift LPG', category='LPG forklifts', unit='liters',
                          qty='100', factor='1.5', month=2, facility=self.facility, description='Warehouse forklifts'),
            self._request(scope='scope2', source='Bought steam', category='Steam', unit='GJ',
                          qty='12.5', factor='60.25', month=4),
            self._request(scope='scope3', source='Courier parcels', category='Last-mile', unit='parcels',
                          qty='2500', factor='0.8', month=7),
        ]
        # Requests the Dashboard does NOT count — none of them may reach a total or an export.
        self._request(scope='scope3', source='Still pending', unit='kg', qty='10', factor=None, status='pending')
        self._request(scope='scope3', source='Was rejected', unit='kg', qty='10', factor=None, status='rejected')
        self._request(scope='scope3', source='Approved, no amount', unit='kg', qty='10', factor='2', calc=False)
        self._request(scope='scope2', source='Facility electricity factor', unit='kWh', qty='1000', factor='0.5',
                      electricity_method='location_based', facility=self.facility)
        self._request(scope='scope3', source='Last year', unit='kg', qty='10', factor='2', year=YEAR - 1)
        self._request(scope='scope3', source='Someone else', unit='kg', qty='10', factor='2', company=self.other)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

    def _request(self, scope, source, unit, qty, factor, month=1, status='approved', year=YEAR, category='x',
                 description='', facility=None, calc=True, electricity_method='', company=None):
        return CustomEmissionRequest.objects.create(
            user=self.user, company=company or self.company, scope=scope, category_name=category,
            source_name=source, description=description, unit=unit, quantity=Decimal(qty), year=year, month=month,
            facility=facility, status=status, electricity_method=electricity_method,
            approved_factor_kg_co2e=Decimal(factor) if factor else None,
            calculated_co2e_kg=(Decimal(qty) * Decimal(factor)) if (factor and calc) else None)

    # ── the three views of one inventory ──────────────────────────────────────
    def _summary(self):
        return self.client.get(f'/api/emissions/summary/?year={YEAR}').data

    def _csv(self, lang='en'):
        res = self.client.get(f'/api/emissions/export-csv/?year={YEAR}&lang={lang}')
        self.assertEqual(res.status_code, 200)
        text = res.content.decode('utf-8').lstrip('﻿')
        rows = list(csv.reader(io.StringIO(text), delimiter=';' if lang == 'tr' else ','))
        return rows[0], rows[1:]

    def _xlsx(self, lang='en'):
        res = self.client.get(f'/api/emissions/export-excel/?year={YEAR}&lang={lang}')
        self.assertEqual(res.status_code, 200)
        rows = [list(r) for r in load_workbook(io.BytesIO(res.content)).active.iter_rows(values_only=True)]
        return rows[0], [r for r in rows[1:] if r and r[0] is not None]

    @staticmethod
    def _num(v):
        return float(str(v).replace(',', '.')) if isinstance(v, str) else float(v)

    def _approved(self, header, rows, lang='en'):
        label = 'Onaylı' if lang == 'tr' else 'Approved'
        st = header.index('Durum' if lang == 'tr' else 'Status')
        return [r for r in rows if str(r[st]).startswith(label)]

    # ── totals ────────────────────────────────────────────────────────────────
    def test_the_dashboard_counts_the_custom_requests(self):
        """Control: the premise of F-2 — the Dashboard total contains the three counted requests."""
        s = self._summary()
        entries_kg = (10 + 4) * 56.211
        custom_kg = 100 * 1.5 + 12.5 * 60.25 + 2500 * 0.8
        self.assertAlmostEqual(s['total_kg'], entries_kg + custom_kg, places=3)
        self.assertAlmostEqual(s['custom_emissions']['approved_total_kg'], custom_kg, places=3)

    def test_each_export_adds_up_to_the_dashboard_total(self):
        total = self._summary()['total_kg']
        for lang in ('en', 'tr'):
            for name, (header, rows) in (('csv', self._csv(lang)), ('xlsx', self._xlsx(lang))):
                kg = header.index('kg CO2e')
                t = header.index('tCO2e')
                approved = self._approved(header, rows, lang)
                self.assertAlmostEqual(sum(self._num(r[kg]) for r in approved), total, places=2, msg=f'{name} {lang} kg')
                self.assertAlmostEqual(sum(self._num(r[t]) for r in approved), total / 1000, places=3, msg=f'{name} {lang} t')

    def test_each_export_adds_up_to_the_dashboard_per_scope(self):
        s = self._summary()
        for name, (header, rows) in (('csv', self._csv()), ('xlsx', self._xlsx())):
            by_scope = {}
            for r in self._approved(header, rows):
                by_scope[r[1]] = by_scope.get(r[1], 0) + self._num(r[header.index('kg CO2e')])
            self.assertAlmostEqual(by_scope.get('Scope 1', 0), s['scope1_kg'], places=2, msg=name)
            self.assertAlmostEqual(by_scope.get('Scope 2', 0), s['scope2_kg'], places=2, msg=name)
            self.assertAlmostEqual(by_scope.get('Scope 3', 0), s['scope3_kg'], places=2, msg=name)

    # ── rows ──────────────────────────────────────────────────────────────────
    def test_every_counted_request_is_one_row_with_its_own_quantity_unit_factor_and_amount(self):
        for name, (header, rows) in (('csv', self._csv()), ('xlsx', self._xlsx())):
            col = {h: i for i, h in enumerate(header)}
            for cr in self.counted:
                [r] = [r for r in rows if r[col['Source']] == cr.source_name]
                self.assertEqual(r[col['Scope']], cr.scope.replace('scope', 'Scope '), name)
                self.assertEqual(r[col['Unit']], cr.unit, name)
                self.assertAlmostEqual(self._num(r[col['Quantity']]), float(cr.quantity), places=4, msg=name)
                self.assertAlmostEqual(self._num(r[col['Factor (kg CO2e/unit)']]), float(cr.approved_factor_kg_co2e), places=6, msg=name)
                self.assertAlmostEqual(self._num(r[col['kg CO2e']]), float(cr.calculated_co2e_kg), places=3, msg=name)
                self.assertAlmostEqual(self._num(r[col['tCO2e']]), float(cr.calculated_co2e_kg) / 1000, places=6, msg=name)
                # the amount is exactly quantity x factor, in the request's own unit
                self.assertAlmostEqual(self._num(r[col['kg CO2e']]),
                                       self._num(r[col['Quantity']]) * self._num(r[col['Factor (kg CO2e/unit)']]), places=3)
                self.assertEqual(r[col['Status']], 'Approved · Custom Request', name)
                self.assertIn(f'#{cr.pk}', r[col['Factor reference']], name)

    def test_a_request_row_carries_its_month_facility_and_description(self):
        header, rows = self._xlsx()
        col = {h: i for i, h in enumerate(header)}
        [r] = [r for r in rows if r[col['Source']] == 'Forklift LPG']
        self.assertEqual((r[col['Month']], r[col['Facility']], r[col['Description']], r[col['Category']]),
                         ('February', 'Gebze', 'Warehouse forklifts', 'LPG forklifts'))

    def test_requests_the_dashboard_does_not_count_are_not_in_the_exports(self):
        not_counted = {'Still pending', 'Was rejected', 'Approved, no amount', 'Facility electricity factor',
                       'Last year', 'Someone else'}
        for header, rows in (self._csv(), self._xlsx()):
            self.assertEqual({r[0] for r in rows} & not_counted, set())

    def test_the_row_count_is_the_entries_plus_the_counted_requests(self):
        for header, rows in (self._csv(), self._xlsx()):
            self.assertEqual(len(rows), 3 + len(self.counted))                  # 2 approved + 1 pending entry

    def test_the_entry_rows_are_unchanged(self):
        header, rows = self._xlsx()
        gas = [r for r in rows if r[0] == 'Natural Gas']
        self.assertEqual(sorted((r[header.index('Quantity')], r[header.index('Status')]) for r in gas),
                         [(4.0, 'Approved'), (10.0, 'Approved'), (99.0, 'Pending')])

    # ── language ──────────────────────────────────────────────────────────────
    def test_the_turkish_exports_label_the_requests_in_turkish(self):
        for header, rows in (self._csv('tr'), self._xlsx('tr')):
            col = {h: i for i, h in enumerate(header)}
            [r] = [r for r in rows if r[col['Kaynak']] == 'Forklift LPG']
            self.assertEqual(r[col['Durum']], 'Onaylı · Özel Talep')
            self.assertEqual((r[col['Kapsam']], r[col['Ay']]), ('Kapsam 1', 'Şubat'))
            self.assertIn('Özel Talep #', r[col['Faktör kaynağı']])

    def test_the_turkish_csv_uses_the_decimal_comma_for_requests_too(self):
        header, rows = self._csv('tr')
        col = {h: i for i, h in enumerate(header)}
        [r] = [r for r in rows if r[col['Kaynak']] == 'Bought steam']
        self.assertEqual((r[col['Miktar']], r[col['Faktör (kg CO2e/birim)']], r[col['kg CO2e']]), ('12,5', '60,25', '753,125'))

    # ── the full backup workbook ──────────────────────────────────────────────
    def test_the_backup_workbook_has_a_sheet_of_the_company_requests(self):
        res = self.client.get('/api/emissions/export-all/?file=xlsx&lang=en')
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(io.BytesIO(res.content))
        self.assertIn('Custom Requests', wb.sheetnames)
        rows = [list(r) for r in wb['Custom Requests'].iter_rows(values_only=True)]
        header, body = rows[0], rows[1:]
        col = {h: i for i, h in enumerate(header)}
        mine = CustomEmissionRequest.objects.filter(company=self.company)
        self.assertEqual(len(body), mine.count())                                # every status, none of another company
        self.assertNotIn('Someone else', {r[col['Source']] for r in body})
        counted = [r for r in body if r[col['Counted in totals']] == 'Yes']
        self.assertEqual({r[col['Source']] for r in counted}, {cr.source_name for cr in self.counted} | {'Last year'})
        self.assertAlmostEqual(sum(float(r[col['kg CO2e']]) for r in counted if r[col['Year']] == YEAR),
                               self._summary()['custom_emissions']['approved_total_kg'], places=3)
